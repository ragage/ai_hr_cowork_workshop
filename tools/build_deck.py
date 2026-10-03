#!/usr/bin/env python
"""Build the instructor deck from the Copilot Cowork scenario PowerPoint template."""
import copy
import os
import re
import uuid

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from content import EXERCISES, GUIDE_CARD

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # repo root (this file is in tools/)
BUILD = os.path.join(ROOT, ".build")
os.makedirs(BUILD, exist_ok=True)
# The source PowerPoint template isn't in the repo; point DECK_TEMPLATE at a local copy.
TEMPLATE = os.environ.get("DECK_TEMPLATE", "")
if not os.path.isfile(TEMPLATE):
    raise SystemExit("Set DECK_TEMPLATE to the path of the source PowerPoint template (.pptx).")
OUT = os.environ.get("DECK_OUT", os.path.join(ROOT, "instructor", "instructor-deck.pptx"))
DECK_SUBJECT = "Instructor deck"
# Download links on the "Workshop kit" slide (both decks); same URLs as the emails in communication/.
KIT_REPO = os.environ.get("KIT_REPO", "https://github.com/cragage_microsoft/ai_hr_cowork_workshop")
HINT_IMG = os.path.join(ROOT, "reference", "media", "download-hint.png")  # from make_download_hint.py


def kit_url(path):
    return f"{KIT_REPO}/raw/main/{path}"
UI_SHOT = os.path.join(ROOT, "tools", "assets", "cowork-home.png")  # Cowork home-page screenshot
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"

# Template palette
INK = "323139"        # scenario-slide text
CARD = "FFFCFD"       # goal/output card
PANEL = "FFF9FD"      # prompt/workflow panel
PILL = "F3E5FD"       # section pills
PLUM = "341434"       # "Why Cowork?" card
PLUM_ICON = "3B1041"
W_TITLE, W_SUB, W_BODY, W_LINE = "242424", "616161", "242424", "D1D1D1"
BLUE, GREEN, PURPLE, RED, OLIVE, GRAY = "0F6CBD", "107C10", "8064A2", "C0504D", "948A54", "8A8886"
SEG_UI, SEG_DISP, SEG_SEMI = "Segoe UI", "Segoe Sans Display", "Segoe Sans Display Semibold"

prs = Presentation(TEMPLATE)
ORIG = list(prs.slides)
N_ORIG = len(ORIG)
S_TITLE, S_WHITE, S_CARD = ORIG[0], ORIG[2], ORIG[8]

# ---------------------------------------------------------------- icons
ICON_GLYPHS = {"web": "\ue774", "onedrive": "\ue753", "sharepoint": "\ue8f1"}
ICON_LABELS = {"m365": "M365 Data", "web": "Web", "onedrive": "OneDrive", "sharepoint": "SharePoint"}
STOPS = [(0.0, (40, 112, 234)), (0.4, (155, 81, 224)), (0.75, (233, 64, 122)), (1.0, (247, 162, 59))]


def _grad(t):
    for (a, ca), (b, cb) in zip(STOPS, STOPS[1:]):
        if a <= t <= b:
            f = (t - a) / (b - a)
            return tuple(int(ca[k] + (cb[k] - ca[k]) * f) for k in range(3))
    return STOPS[-1][1]


def make_icon(key, size=256):
    path = os.path.join(BUILD, f"icon_{key}.png")
    font = ImageFont.truetype(r"C:\Windows\Fonts\SegoeIcons.ttf", int(size * 0.8))
    mask = Image.new("L", (size, size), 0)
    d = ImageDraw.Draw(mask)
    bb = d.textbbox((0, 0), ICON_GLYPHS[key], font=font)
    d.text(((size - (bb[2] - bb[0])) / 2 - bb[0], (size - (bb[3] - bb[1])) / 2 - bb[1]),
           ICON_GLYPHS[key], font=font, fill=255)
    grad = Image.new("RGBA", (size, size))
    px = grad.load()
    for yy in range(size):
        for xx in range(size):
            px[xx, yy] = _grad((xx + yy) / (2 * (size - 1))) + (255,)
    out = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    out.paste(grad, (0, 0), mask)
    out.save(path)
    return path


ICON_PATHS = {k: make_icon(k) for k in ICON_GLYPHS}


# ---------------------------------------------------------------- slide cloning
def clone_slide(src):
    new = prs.slides.add_slide(src.slide_layout)
    rid_map = {}
    for rId, rel in list(src.part.rels.items()):
        if rel.reltype in (RT.NOTES_SLIDE, RT.SLIDE_LAYOUT):
            continue
        if rel.is_external:
            nr = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype.endswith("/themeOverride"):
            tp = rel.target_part
            pn = prs.part.package.next_partname("/ppt/theme/themeOverride%d.xml")
            nr = new.part.relate_to(Part(pn, tp.content_type, prs.part.package, tp.blob), rel.reltype)
        else:
            nr = new.part.relate_to(rel.target_part, rel.reltype)
        rid_map[rId] = nr
    ne, se = new._element, src._element
    # Keep the existing cSld/spTree elements (python-pptx caches them); swap their contents.
    n_cSld, s_cSld = ne.find(qn("p:cSld")), se.find(qn("p:cSld"))
    n_tree = n_cSld.find(qn("p:spTree"))
    for child in list(n_cSld):
        if child is not n_tree:
            n_cSld.remove(child)
    for child in list(n_tree):
        n_tree.remove(child)
    for k, v in s_cSld.attrib.items():
        n_cSld.set(k, v)
    before = True
    for child in s_cSld:
        if child.tag == qn("p:spTree"):
            for c in child:
                n_tree.append(copy.deepcopy(c))
            before = False
        elif before:
            n_tree.addprevious(copy.deepcopy(child))
        else:
            n_cSld.append(copy.deepcopy(child))
    s_cm, n_cm = se.find(qn("p:clrMapOvr")), ne.find(qn("p:clrMapOvr"))
    if s_cm is not None:
        if n_cm is not None:
            ne.replace(n_cm, copy.deepcopy(s_cm))
        else:
            n_cSld.addnext(copy.deepcopy(s_cm))
    for el in ne.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith("{%s}" % R_NS) and v in rid_map:
                el.set(k, rid_map[v])
    return new


def by_name(slide):
    found = {}

    def walk(shapes):
        for sh in shapes:
            found.setdefault(sh.name, sh)
            if sh.shape_type == 6:
                walk(sh.shapes)
    walk(slide.shapes)
    return found


def keep_only(slide, names):
    for sh in list(slide.shapes):
        if sh.name not in names:
            sh._element.getparent().remove(sh._element)


def dedupe_ids(slide):
    seen, els = set(), list(slide._element.iter(qn("p:cNvPr")))
    top = max(int(e.get("id")) for e in els)
    for e in els:
        i = int(e.get("id"))
        if i in seen:
            top += 1
            e.set("id", str(top))
        seen.add(int(e.get("id")))


def set_first_run(shape, text, size=None):
    tf = shape.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    runs = tf.paragraphs[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if size:
        runs[0].font.size = Pt(size)


def set_lines(shape, lines, size=None):
    """Title text on deliberate lines, separated by soft line breaks."""
    set_first_run(shape, lines[0], size)
    r0 = shape.text_frame.paragraphs[0]._p.findall(qn("a:r"))[0]
    prev = r0
    for ln in lines[1:]:
        br = etree.Element(qn("a:br"))
        rpr = r0.find(qn("a:rPr"))
        if rpr is not None:
            br.append(copy.deepcopy(rpr))
        prev.addnext(br)
        r = copy.deepcopy(r0)
        r.find(qn("a:t")).text = ln
        br.addnext(r)
        prev = r


# ---------------------------------------------------------------- paragraph cloning (template text)
def protos(shape):
    return [copy.deepcopy(p) for p in shape.text_frame._txBody.findall(qn("a:p"))]


def build_p(proto, segs):
    """segs: (text, proto_run_index[, bold])"""
    p = copy.deepcopy(proto)
    rprs = [copy.deepcopy(r.find(qn("a:rPr"))) for r in p.findall(qn("a:r"))]
    for child in list(p):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p.remove(child)
    end = p.find(qn("a:endParaRPr"))
    for seg in segs:
        text, ri = seg[0], seg[1]
        r = etree.Element(qn("a:r"))
        if rprs and ri < len(rprs) and rprs[ri] is not None:
            rp = copy.deepcopy(rprs[ri])
            if len(seg) > 2:
                rp.set("b", "1" if seg[2] else "0")
            r.append(rp)
        etree.SubElement(r, qn("a:t")).text = text
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)
    return p


def set_paras(shape, ps):
    tb = shape.text_frame._txBody
    for p in tb.findall(qn("a:p")):
        tb.remove(p)
    for p in ps:
        tb.append(p)


def set_sizes(shape, sz):
    for el in shape._element.iter(qn("a:rPr"), qn("a:endParaRPr")):
        el.set("sz", str(int(sz * 100)))


# ---------------------------------------------------------------- native drawing helpers
def _shadow(sp, alpha="16000"):
    spPr = sp._element.spPr
    ef = spPr.find(qn("a:effectLst"))
    if ef is None:
        ef = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(ef, qn("a:outerShdw"), blurRad="76200", dist="19050", dir="5400000",
                          algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
    etree.SubElement(clr, qn("a:alpha"), val=alpha)


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, shadow=False):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(0.75)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if radius is not None:
        sp.adjustments[0] = radius
    if shadow:
        _shadow(sp)
    sp.text_frame.text = ""
    return sp


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, font=SEG_UI,
         color=W_BODY, size=12, space_after=4, line_spacing=None):
    """paras: list of paragraphs; each is a str, or dict(runs=[(text, {opts})], bullet='num'|'dot')."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        if isinstance(para, str):
            para = {"runs": [(para, {})]}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        p.space_after = Pt(para.get("space_after", space_after))
        if line_spacing:
            p.line_spacing = line_spacing
        for t, o in para["runs"]:
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = RGBColor.from_string(o.get("color", color))
        bullet = para.get("bullet")
        if bullet:
            pPr = p._p.get_or_add_pPr()
            ind = Emu(Inches(0.26 if bullet == "num" else 0.2))
            pPr.set("marL", str(ind))
            pPr.set("indent", str(-ind))
            if bullet == "num":
                etree.SubElement(pPr, qn("a:buFont"), typeface="+mj-lt")
                etree.SubElement(pPr, qn("a:buAutoNum"), type="arabicPeriod")
            else:
                etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
                etree.SubElement(pPr, qn("a:buChar"), char="\u2022")
    return tb


HLINK_CLR_URI = "{A12FA001-AC4F-418D-AE19-62706E023703}"
AHYP = "http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor"


def link_run(r, url):
    """Hyperlink a run and keep its own text color (instead of the theme's hyperlink blue)."""
    r.hyperlink.address = url
    h = r._r.find(qn("a:rPr")).find(qn("a:hlinkClick"))
    ext = etree.SubElement(etree.SubElement(h, qn("a:extLst")), qn("a:ext"), uri=HLINK_CLR_URI)
    etree.SubElement(ext, "{%s}hlinkClr" % AHYP, nsmap={"ahyp": AHYP}, val="tx")


def link_list(slide, x, y, w, h, items, accent, size=16):
    """items: (label, url, description) — the label is a hyperlink, the description follows it."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (label, url, desc) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(11)
        r = p.add_run()
        r.text = label
        r.font.name, r.font.size, r.font.bold, r.font.underline = SEG_UI, Pt(size), True, True
        r.font.color.rgb = RGBColor.from_string(accent)
        link_run(r, url)
        r2 = p.add_run()
        r2.text = "  \u00b7  " + desc
        r2.font.name, r2.font.size = SEG_UI, Pt(size - 2)
        r2.font.color.rgb = RGBColor.from_string(W_SUB)
    return tb


def rich(md, size, color, font=SEG_UI, bold_font=None):
    """'**bold** normal' → run list."""
    runs, parts = [], md.split("**")
    for i, part in enumerate(parts):
        if part:
            o = {"size": size, "color": color, "font": bold_font if (i % 2 and bold_font) else font}
            if i % 2:
                o["bold"] = True
            runs.append((part, o))
    return runs


def pill(slide, cx, y, label, w=1.68, h=0.3):
    sp = box(slide, cx - w / 2, y, w, h, PILL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(slide, cx - w / 2, y, w, h, [{"runs": [(label, {"font": SEG_SEMI, "size": 12, "color": INK})],
                                       "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    return sp


NOTES_BODY = (
    '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<p:nvSpPr><p:cNvPr id="%d" name="Notes Placeholder 2"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
    '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="685800" y="4343400"/><a:ext cx="5486400" cy="4114800"/></a:xfrm></p:spPr>'
    '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>')


def notes(slide, t):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is None:
        tree = ns.shapes._spTree
        nid = max([int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))] + [1]) + 1
        tree.append(etree.fromstring(NOTES_BODY % nid))
        tf = ns.notes_text_frame
    tf.text = t


# ---------------------------------------------------------------- prompt elements
# Goal · Source · Expectations · Constraints: same colors as the Word workbook (make_word_docs.py).
PROMPT_EL = {"G": ("Goal", "DCEBFF", "1F4E99", "What outcome do I want, for whom, and why?"),
             "S": ("Source", "DFF5E1", "1E6B32", "Which files, sites, or data should Cowork use?"),
             "E": ("Expectations", "FFE9CC", "8A4B00", "What should the result look like: format, length, tone?"),
             "C": ("Constraints", "EDE3FA", "5B2C91", "What must Cowork not do: guess, invent, send?")}


RPR_BEFORE_HL = {qn(t) for t in ("a:ln", "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill",
                                   "a:grpFill", "a:effectLst", "a:effectDag")}


def highlight(r, fill):
    """Shade an a:r element like a highlighter (PowerPoint for Microsoft 365), keeping rPr schema order."""
    rPr = r.find(qn("a:rPr"))
    if rPr is None:
        rPr = etree.Element(qn("a:rPr"))
        r.insert(0, rPr)
    hl = etree.Element(qn("a:highlight"))
    etree.SubElement(hl, qn("a:srgbClr"), val=fill)
    before = [c for c in rPr if c.tag in RPR_BEFORE_HL]
    if before:
        before[-1].addnext(hl)
    else:
        rPr.insert(0, hl)


PROMPT_TAG = re.compile(r"\{([gsec])\}(.*?)\{/\1\}")


def prompt_segs(line):
    """'{g}text{/g} plain' -> [(text, 'G'), (' plain', None)] (content.py marks prompt elements this way)."""
    out, pos = [], 0
    for m in PROMPT_TAG.finditer(line):
        if m.start() > pos:
            out.append((line[pos:m.start()], None))
        out.append((m.group(2), m.group(1).upper()))
        pos = m.end()
    if pos < len(line):
        out.append((line[pos:], None))
    return out



# ---------------------------------------------------------------- template-based slide builders
def white_slide(title, subtitle):
    s = clone_slide(S_WHITE)
    keep_only(s, {"TextBox 1", "TextBox 2"})
    d = by_name(s)
    set_first_run(d["TextBox 1"], title)
    set_first_run(d["TextBox 2"], subtitle)
    return s


def wcard(slide, x, y, w, h, accent, heading, paras, hsize=14, bsize=12, bar="top"):
    box(slide, x, y, w, h, "FFFFFF", line=W_LINE, shadow=True)
    if bar == "top":
        box(slide, x, y, w, 0.045, accent)
        text(slide, x + 0.16, y + 0.2, w - 0.32, 0.4,
             [{"runs": [(heading, {"size": hsize, "bold": True, "color": accent})]}])
        text(slide, x + 0.16, y + 0.68, w - 0.32, h - 0.8, paras, size=bsize, space_after=5)
    else:
        box(slide, x, y, 0.06, h, accent)


def band(slide, x, y, w, h, md, fill=PILL, color=INK, size=13):
    box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(slide, x + 0.3, y, w - 0.6, h, [{"runs": rich(md, size, color)}], anchor=MSO_ANCHOR.MIDDLE,
         line_spacing=1.05)


def bullets(items, size=12, color=W_BODY, kind="dot"):
    return [{"runs": rich(i, size, color), "bullet": kind} for i in items]


def scenario_card(ex):
    s = clone_slide(S_CARD)
    d = by_name(s)
    set_first_run(d["Title 3"], ex["title"])
    set_first_run(d["NavPill_SitesPages"], ex["pill"])
    set_first_run(d["TextBox 30"], ex.get("tag_label", "Function"))
    set_first_run(d["TextBox 49"], ex["function"])
    for name, label, key in (("Rectangle: Rounded Corners 32", "Goal", "goal"),
                             ("Rectangle: Rounded Corners 33", "Output", "output")):
        set_paras(d[name], [build_p(protos(d[name])[0], [(label, 0), (": ", 1), (ex[key], 2)])])
    why = d["Rectangle: Rounded Corners 97"]
    set_paras(why, [build_p(protos(why)[0], [("Why %s?: " % ex.get("why_label", "Cowork"), 0), (ex["why"], 3)])])
    pr = d["Text 25"]
    pp = protos(pr)
    p_plain, p_empty, p_bullet = pp[0], pp[1], pp[5]
    new = []
    for line in ex["prompt"]:
        if line == "":
            new.append(copy.deepcopy(p_empty))
        elif line.startswith("- "):
            new.append(build_p(p_bullet, [(line[2:], 0)]))
        elif line.startswith("## "):
            new.append(build_p(p_plain, [(line[3:], 0, True)]))
        else:
            segs = prompt_segs(line)
            para = build_p(p_plain, [(t, 0) for t, _ in segs])
            for r, (_, key) in zip(para.findall(qn("a:r")), segs):
                if key:
                    highlight(r, PROMPT_EL[key][1])
            new.append(para)
    set_paras(pr, new)
    if ex.get("prompt_size"):
        set_sizes(pr, ex["prompt_size"])
    wf = d["TextBox 56"]
    wproto = protos(wf)[0]
    set_paras(wf, [build_p(wproto, [(lbl, 0), (" \u2192 " + txt, 1)]) for lbl, txt in ex["workflow"]])
    g = d["Group 12"]
    orig = copy.deepcopy(g._element)
    x0 = int(g._element.find(qn("p:grpSpPr")).find(qn("a:xfrm")).find(qn("a:off")).get("x"))
    prev = g._element
    for i, src in enumerate(ex["sources"]):
        ge = g._element if i == 0 else copy.deepcopy(orig)
        if i:
            prev.addnext(ge)
            prev = ge
        ge.find(qn("p:grpSpPr")).find(qn("a:xfrm")).find(qn("a:off")).set("x", str(x0 + int(Inches(1.0)) * i))
        for t in ge.iter(qn("a:t")):
            t.text = ICON_LABELS[src]
        if src != "m365":
            blip = ge.find(".//" + qn("a:blip"))
            _, rId = s.part.get_or_add_image_part(ICON_PATHS[src])
            blip.set(qn("r:embed"), rId)
            ext = blip.find(qn("a:extLst"))
            if ext is not None:
                blip.remove(ext)
    dedupe_ids(s)
    notes(s, ex["notes_card"])
    return s


def hands_on(ex):
    s = clone_slide(S_CARD)
    keep_only(s, {"Title 3", "NavPill_SitesPages", "TextBox 29", "TextBox 30", "TextBox 49"})
    d = by_name(s)
    set_first_run(d["Title 3"], "Hands-on \u00b7 " + ex["short"])
    set_first_run(d["NavPill_SitesPages"], ex["pill"])
    set_first_run(d["TextBox 30"], "Time")
    set_first_run(d["TextBox 49"], ex["minutes"])
    # Steps panel
    box(s, 0.46, 1.3, 7.84, 3.95, PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    pill(s, 4.38, 1.15, "Steps")
    text(s, 0.85, 1.72, 7.1, 3.45, bullets(ex["steps"], size=15, color=INK, kind="num"), font=SEG_DISP,
         space_after=9, line_spacing=1.05)
    # Discuss panel
    box(s, 0.46, 5.55, 7.84, 1.5, PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    pill(s, 4.38, 5.4, "Discuss")
    text(s, 0.85, 5.8, 7.1, 1.1, [{"runs": rich(ex["discuss"], 14, INK, font=SEG_DISP)}],
         anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    # Watch-for panel
    box(s, 8.49, 1.3, 4.62, 2.4, PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    pill(s, 10.8, 1.15, "Watch for")
    text(s, 8.75, 1.72, 4.15, 1.9, bullets(ex["watch"], size=11.5, color=INK), font=SEG_DISP, space_after=4)
    # Checkpoint (plum, echoing "Why Cowork?")
    box(s, 8.49, 3.95, 4.62, 1.85, PLUM, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    box(s, 8.25, 3.72, 0.44, 0.44, "FFFFFF", shape=MSO_SHAPE.OVAL, shadow=True)
    text(s, 8.25, 3.72, 0.44, 0.44, [{"runs": [("\u2713", {"size": 16, "bold": True, "color": PLUM_ICON})],
                                      "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.75, 4.1, 4.15, 1.6, [{"runs": [("Checkpoint: ", {"size": 12, "bold": True, "color": "FFFFFF", "font": SEG_SEMI}),
                                             (ex["checkpoint"], {"size": 12, "color": "FFFFFF", "font": SEG_DISP})]}],
         anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    # Stretch
    box(s, 8.49, 6.0, 4.62, 1.05, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, 8.75, 6.0, 4.15, 1.05, [{"runs": [("Stretch: ", {"size": 11, "bold": True, "color": INK, "font": SEG_SEMI}),
                                              (ex["stretch"], {"size": 11, "color": INK, "font": SEG_DISP})]}],
         anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.03)
    notes(s, ex["notes_hands"])
    return s


def gradient_blank():
    s = clone_slide(S_CARD)
    keep_only(s, set())
    return s


def title_slide(title, subtitle, title_size=None, lines=None):
    s = clone_slide(S_TITLE)
    d = by_name(s)
    if lines:
        set_lines(d["Title 1"], lines, title_size)
        d["Title 1"].top, d["Title 1"].height = Inches(2.35), Inches(1.5)
    else:
        set_first_run(d["Title 1"], title, title_size)
    set_first_run(d["Text Placeholder 4"], subtitle)
    return s


def break_slide(next_up):
    s = gradient_blank()
    text(s, 1.0, 2.35, 11.3, 1.3, [{"runs": [("\u2615  Break \u00b7 10 minutes", {"font": SEG_DISP, "size": 54, "color": INK})],
                                    "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.5, 3.75, 10.3, 0.6, [{"runs": [("Stretch, refill, and note one task you'd hand to Cowork tomorrow.",
                                              {"font": SEG_DISP, "size": 18, "color": INK})], "align": PP_ALIGN.CENTER}])
    box(s, 3.9, 4.75, 5.53, 0.55, PILL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, 3.9, 4.75, 5.53, 0.55, [{"runs": [("Next up: " + next_up, {"font": SEG_SEMI, "size": 15, "color": INK})],
                                     "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    return s


# ================================================================ BUILD
SECTION_MARKS = []  # (section name, 1-based final slide number where it starts)
DIVIDER_ALT = {1: "Illustration of a bar chart, for the executive command center dashboard",
               2: "Illustration of a magnifying glass, for Deep Research",
               3: "Illustration of a globe, for navigating websites with Cowork's browser",
               4: "Illustration of a hammer and wrench, for building a custom skill",
               5: "Illustration of a clipboard, for recruiting and reporting",
               6: "Illustration of a party popper, for welcoming a new hire",
               7: "Illustration of an alarm clock, for scheduled automations",
               8: "Illustration of a robot, for the HR Policy Agent"}


def mark_section(name):
    SECTION_MARKS.append((name, len(prs.slides._sldIdLst) - N_ORIG + 1))


def exercise_divider(ex):
    """Full-bleed divider before each exercise: big number, title, time, and an 8-step progress tracker."""
    import re as _re
    num, total_ex = ex["num"], len(EXERCISES)
    m = _re.search(r"\((\d:\d\d)-(\d:\d\d), (\d+) min\)", ex["notes_card"])
    clock = f"{m.group(1)}\u2013{m.group(2)}" if m else ""
    s = gradient_blank()
    text(s, 0.95, 1.45, 3.0, 0.4, [{"runs": [("EXERCISE", {"font": SEG_SEMI, "size": 16, "color": W_SUB})]}])
    text(s, 0.85, 1.8, 3.2, 2.0, [{"runs": [(f"{num:02d}", {"font": SEG_DISP, "size": 120, "bold": True,
                                                             "color": PURPLE})]}], anchor=MSO_ANCHOR.TOP)
    title = ex["title"]
    text(s, 4.1, 1.55, 5.95, 1.65, [{"runs": [(title, {"font": SEG_DISP, "size": 38 if len(title) < 26 else 32,
                                                       "color": INK})]}], anchor=MSO_ANCHOR.BOTTOM, line_spacing=0.95)
    text(s, 4.1, 3.3, 5.95, 1.1, [{"runs": [(ex["goal"], {"font": SEG_DISP, "size": 16, "color": W_SUB})]}],
         line_spacing=1.05)
    # Picture that describes the exercise (Fluent Emoji 3D, MIT license), in a soft white circle
    art = os.path.join(os.path.dirname(os.path.abspath(__file__)), "art", f"ex{num}.png")
    if os.path.exists(art):
        box(s, 10.3, 1.4, 2.4, 2.4, "FFFFFF", line=W_LINE, shadow=True, shape=MSO_SHAPE.OVAL)
        pic = s.shapes.add_picture(art, Inches(10.7), Inches(1.8), width=Inches(1.6))
        pic.name = f"Picture Divider {num}"
        pic._element.nvPicPr.cNvPr.set("descr", DIVIDER_ALT.get(num, title))
    chips = [ex["minutes"] + ("  \u00b7  " + clock if clock else ""), ex["function"]]
    x = 4.1
    for c in chips:
        w = 0.4 + 0.105 * len(c)
        box(s, x, 4.45, w, 0.48, PILL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        text(s, x, 4.45, w, 0.48, [{"runs": [(c, {"font": SEG_SEMI, "size": 14, "color": INK})],
                                    "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
        x += w + 0.2
    # progress tracker
    d, gap = 0.56, 0.36
    x0 = (13.333 - (total_ex * d + (total_ex - 1) * gap)) / 2
    for i in range(1, total_ex + 1):
        cx = x0 + (i - 1) * (d + gap)
        if i == num:
            fill, line, col = PURPLE, None, "FFFFFF"
        elif i < num:
            fill, line, col = PILL, None, INK
        else:
            fill, line, col = "FFFFFF", W_LINE, W_SUB
        box(s, cx, 5.75, d, d, fill, line=line, shape=MSO_SHAPE.OVAL)
        text(s, cx, 5.75, d, d, [{"runs": [(str(i), {"font": SEG_SEMI, "size": 14, "bold": i == num, "color": col})],
                                  "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
        if i < total_ex:
            box(s, cx + d + 0.05, 5.75 + d / 2 - 0.01, gap - 0.1, 0.02, W_LINE)
    text(s, 0.5, 6.45, 12.33, 0.35, [{"runs": [(f"Exercise {num} of {total_ex}", {"font": SEG_UI, "size": 12,
                                                                                   "color": W_SUB})],
                                       "align": PP_ALIGN.CENTER}])
    notes(s, f"DIVIDER \u2014 EXERCISE {num}: {title} ({ex['minutes']}{', ' + clock if clock else ''}). "
             "Pause here: check the room is ready (anyone still finishing the last exercise?), then introduce the "
             "scenario card on the next slide.")
    return s


# ---------------------------------------------------------------- slides shared with build_overview.py
def agenda_slide():
    """Agenda slide: the 4-hour run of show (shared by the instructor and overview decks)."""
    s = white_slide("Agenda \u2014 4 hours, hands-on",
                    "Framing and a UI tour, then eight scenario-based exercises with two breaks.")
    LEFT = [("10 min", "Welcome & context", "What Cowork is, HR value, the approval model", BLUE),
            ("10 min", "Copilot vs. Cowork", "The difference and when to use which", BLUE),
            ("20 min", "Cowork UI walkthrough + setup", "New task \u00b7 My tasks \u00b7 Automations \u00b7 Customize", BLUE),
            ("25 min", "Ex 1 \u00b7 Executive Command Center", "Interactive HTML dashboard \u00b7 approvals one at a time", PURPLE),
            ("20 min", "Ex 2 \u00b7 Deep Research", "Cited briefing + Word/Excel scorecard", GREEN),
            ("10 min", "Break", "", GRAY),
            ("25 min", "Ex 3 \u00b7 Navigate websites", "Cowork drives Edge: dol.gov + lni.wa.gov", GREEN)]
    RIGHT = [("30 min", "Ex 4 \u00b7 Build a custom skill", "Centerpiece: HR Policy Answer", GREEN),
             ("15 min", "Ex 5 \u00b7 Recruiting + reporting", "Inclusive job posting + ticket summary", GREEN),
             ("10 min", "Break", "", GRAY),
             ("20 min", "Ex 6 \u00b7 Onboarding pack", "PowerPoint + Scheduling + Communications", GREEN),
             ("15 min", "Ex 7 \u00b7 Automate & share", "Automations, Daily Briefing, sharing", GREEN),
             ("5 min", "Spotlight \u00b7 HR plugins", "Customize \u2192 Plugins, talk only", OLIVE),
             ("20 min", "Ex 8 \u00b7 Agent Builder", "Non-Cowork: build an HR Policy Agent", RED),
             ("5 min", "Wrap-up", "Three tools and next steps", BLUE)]
    for col, rows in enumerate((LEFT, RIGHT)):
        x = 0.55 + col * 6.2
        for i, (t, name, sub, acc) in enumerate(rows):
            y = 1.6 + i * 0.68
            rh = 0.6
            box(s, x, y, 6.0, rh, "FFFFFF", line=W_LINE, shadow=True)
            box(s, x, y, 0.06, rh, acc)
            text(s, x + 0.2, y, 0.9, rh, [{"runs": [(t, {"size": 12, "bold": True, "color": acc})]}], anchor=MSO_ANCHOR.MIDDLE)
            body = [{"runs": [(name, {"size": 13, "bold": True, "color": W_TITLE})], "space_after": 0}]
            if sub:
                body.append({"runs": [(sub, {"size": 10.5, "color": W_SUB})], "space_after": 0})
            text(s, x + 1.2, y, 4.7, rh, body, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "AGENDA. Framing (what/why + Copilot vs Cowork), a UI tour + setup, then EIGHT exercises with two "
             "breaks. Each exercise opens with a scenario card, then a hands-on slide with "
             "steps and a checkpoint. Ex 1 is the Executive Command Center; Ex 3 has Cowork drive a browser; Ex 4 (custom "
             "skill) is the centerpiece; Ex 8 steps OUTSIDE Cowork into Agent Builder.")
    return s


def objectives_slide():
    """Six learning objectives (shared by the instructor and overview decks)."""
    s = white_slide("What you\u2019ll be able to do by the end", "Six learning objectives \u2014 we\u2019ll check them at wrap-up.")
    OBJ = [("Choose the right tool", "Copilot Chat, **Cowork**, or **Agent Builder** \u2014 for the task in front of you.", BLUE),
           ("Delegate safely", "Describe an outcome; approve Cowork\u2019s actions **one at a time**.", PURPLE),
           ("Ground in real content", "Attach files, research and browse the web **with citations**, and check the result.", GREEN),
           ("Package repeatable work", "Build a **custom skill**; schedule recurring work with **Automations**.", RED),
           ("Build a no-code agent", "Answers from HR documents; **declines** what it shouldn\u2019t answer.", OLIVE),
           ("Apply HR guardrails", "Fictional data, **cite and confirm**, people decisions stay with people.", BLUE)]
    for i, (h, sub, acc) in enumerate(OBJ):
        x, y = 0.55 + (i % 2) * 6.2, 1.6 + (i // 2) * 1.75
        box(s, x, y, 6.03, 1.55, "FFFFFF", line=W_LINE, shadow=True)
        box(s, x, y, 0.06, 1.55, acc)
        text(s, x + 0.3, y, 0.7, 1.55, [{"runs": [(str(i + 1), {"size": 32, "bold": True, "color": acc})]}], anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 1.1, y + 0.1, 4.75, 1.35, [{"runs": [(h, {"size": 16, "bold": True, "color": W_TITLE})], "space_after": 4},
                                               {"runs": rich(sub, 13, W_SUB)}], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "LEARNING OBJECTIVES (part of Welcome, 0:00-0:10). Read the six out loud and tell the room you'll come "
             "back to them at wrap-up with a quick knowledge check. They mirror the README and workbook.")
    return s


def kit_links_slide():
    """Download links for participants and instructors (shared by the instructor and overview decks)."""
    s = white_slide("Workshop kit \u2014 download links",
                    "Everything participants and instructors need. Select a link to download it.")
    KIT_PARTICIPANT = [("Participant workbook", kit_url("participant/participant-workbook.docx"), "setup + all 8 exercises"),
                       ("Sample data (zip)", kit_url("participant/zava-sample-knowledge.zip"), "the six Zava files"),
                       ("Quick-reference card", kit_url("participant/quick-reference-card.docx"), "one-page handout"),
                       ("Prompt library", kit_url("reference/04-prompt-library.docx"), "copy-paste HR prompts"),
                       ("Responsible use", kit_url("reference/06-responsible-use.docx"), "data-handling rules"),
                       ("After the workshop", kit_url("participant/after-the-workshop.docx"), "30-day adoption plan")]
    KIT_INSTRUCTOR = [("Instructor deck", kit_url("instructor/instructor-deck.pptx"), "slides + speaker notes"),
                      ("Facilitator guide", kit_url("instructor/facilitator-guide.docx"), "run sheet + triage"),
                      ("Answer key", kit_url("instructor/facilitator-answer-key.docx"), "expected results"),
                      ("Readiness checklist", kit_url("instructor/readiness-checklist.docx"), "prep timeline + Plan B"),
                      ("Seed content", kit_url("instructor/seed-content.docx"), "mail, meetings, chat for Ex 1"),
                      ("Communication kit", f"{KIT_REPO}/tree/main/communication", "overview deck + emails")]
    for col, (head, items, acc) in enumerate((("FOR PARTICIPANTS", KIT_PARTICIPANT, BLUE),
                                              ("FOR INSTRUCTORS", KIT_INSTRUCTOR, PURPLE))):
        x = 0.55 + col * 6.25
        wcard(s, x, 1.6, 5.98, 3.2, acc, head, [])
        link_list(s, x + 0.2, 2.3, 5.6, 2.45, items, acc)
    box(s, 0.55, 5.0, 12.23, 1.85, PLUM, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    hint = s.shapes.add_picture(HINT_IMG, Inches(7.08), Inches(5.1), height=Inches(1.65))
    hint.name = "Picture DownloadHint"
    hint._element.nvPicPr.cNvPr.set("descr", "Tip: on a GitHub file page, select the Download icon (Download raw "
                                               "file) at the top right of the file.")
    tb = s.shapes.add_textbox(Inches(0.85), Inches(5.0), Inches(5.95), Inches(1.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    for t, o in (("Everything at once: ", {"bold": True}), ("download the ", {}),
                 ("whole kit as a zip", {"bold": True, "link": f"{KIT_REPO}/archive/refs/heads/main.zip"}),
                 (", or browse it on ", {}), ("GitHub", {"bold": True, "link": KIT_REPO}),
                 (". Link opens on GitHub? Select the ", {}), ("Download", {"bold": True}),
                 (" icon at the top right of the file. No access? Use the shared Teams/SharePoint folder.", {})):
        r = p.add_run()
        r.text = t
        r.font.name, r.font.size, r.font.bold = SEG_UI, Pt(15), o.get("bold", False)
        r.font.color.rgb = RGBColor.from_string("FFFFFF")
        if o.get("link"):
            r.font.underline = True
            link_run(r, o["link"])
    notes(s, "WORKSHOP KIT (part of setup, 0:20-0:40). Leave this up while people sign in. Participants need the "
             "workbook and the sample-data zip today; the quick-reference card, prompt library, responsible-use guide, "
             "and after-the-workshop pack are take-aways. The right column is for you and the proctors. The links point "
             "to the kit's GitHub repo, which is private: before the session, give attendees access or post the files in "
             "the shared Teams/SharePoint folder and send that link instead (communication/ has the training overview "
             "deck and the invitation emails).")
    return s


# 1 — Title
mark_section("Welcome & orientation")
s = title_slide("Getting Things Done with Copilot Cowork for HR Tasks",
                "Instructor-led, hands-on workshop \u00b7 4 hours \u00b7 ~25 attendees", title_size=40,
                lines=["Getting Things Done with Copilot Cowork", "for HR Tasks"])
notes(s, "WELCOME (0:00-0:10). Introduce yourself and the goal: by the end, every attendee knows when to "
         "use Copilot Chat vs. Cowork and has built an executive command center, researched the web, "
         "had Cowork navigate websites in its browser, built a custom skill, an onboarding pack and an automation, and finished "
         "with a no-code AGENT. 4 hours, two breaks. Tenant note: ATTENDEES share one tenant; YOU demo "
         "from a SEPARATE tenant, so your screen may differ. Golden rule all day: every Cowork output is "
         "a DRAFT; Cowork pauses at checkpoints.")

# 2 — Agenda
agenda_slide()

# 2b — Learning objectives
objectives_slide()

# 3 — Copilot vs Cowork
s = white_slide("Copilot vs. Cowork \u2014 what\u2019s the difference?",
                "Start here: the mental model for the whole day.")
wcard(s, 0.55, 1.6, 5.98, 3.3, BLUE, "COPILOT (CHAT) = YOUR AI ASSISTANT",
      bullets(["**You ask, it answers** \u2014 you drive each step",
               "Best for quick answers, drafting, summarizing, rewriting",
               "Output: a response in the chat",
               "HR example: \u201cSummarize this policy in 3 bullets.\u201d"], size=15), hsize=15)
wcard(s, 6.8, 1.6, 5.98, 3.3, PURPLE, "COWORK = YOUR AI COWORKER",
      bullets(["**You describe an outcome** \u2014 it plans and does the multi-step work",
               "Best for tasks that end in a doc, email, deck, report, or schedule",
               "Output: finished artifacts \u2014 and it can run on a schedule",
               "HR example: \u201cDraft welcome emails for all 4 new hires and save them as drafts.\u201d"], size=15), hsize=15)
band(s, 0.55, 5.2, 12.23, 1.55,
     "**Rule of thumb:** several steps that end in a file, email, schedule, or report \u2192 **Cowork**. A quick "
     "answer or rewrite \u2192 **Copilot Chat**. Both are grounded in your work through Work IQ and respect the "
     "same permissions. We close with a third tool \u2014 **Agent Builder** \u2014 in Exercise 8.")
notes(s, "COPILOT vs COWORK (0:10-0:20). Copilot (Chat) = assistant: you ask, it answers, you drive. Cowork = "
         "coworker: you describe an outcome, it does the multi-step work and returns a deliverable, pausing "
         "for approval. Give one HR example of each and land the rule of thumb. Cowork isn't a replacement "
         "for Chat — it's the 'do it for me' mode of the same Copilot family. Detail: "
         "reference/00-copilot-vs-cowork.md.")

# 4 — What is Cowork
s = white_slide("What is Copilot Cowork?",
                "You describe the outcome \u2014 Cowork plans and delivers, grounded in your work via Work IQ.")
for i, (h, sub, eg, acc) in enumerate([("1 \u00b7 PLANS", "the steps needed to reach your outcome", "e.g., gather \u2192 draft \u2192 schedule", BLUE),
                                       ("2 \u00b7 PICKS", "the right skills and apps for each step", "e.g., Email, Word, Calendar", PURPLE),
                                       ("3 \u00b7 PAUSES", "for your approval at key checkpoints", "e.g., before it sends or deletes", GREEN),
                                       ("4 \u00b7 RETURNS", "finished artifacts, ready for your review", "e.g., a deck, report, or dashboard", RED)]):
    wcard(s, 0.55 + i * 3.1, 1.6, 2.93, 2.95, acc, h,
          [{"runs": rich(sub, 16, W_BODY), "space_after": 14},
           {"runs": [(eg, {"size": 13, "italic": True, "color": W_SUB})]}], hsize=16)
band(s, 0.55, 4.9, 12.23, 1.85,
     "**Why HR cares \u2014 the checkpoint model:** nothing irreversible (like sending an email) happens without "
     "your approval. Every output is a draft to review, so HR stays accountable. You\u2019ll "
     "approve actions one at a time all day.", fill=PLUM, color="FFFFFF", size=16)
notes(s, "WHAT IS COWORK. Describe the OUTCOME, not the steps. Cowork plans -> picks skills -> pauses for "
         "approval -> returns artifacts, grounded via Work IQ (permission-aware). Hammer the checkpoint "
         "model; attendees meet approval dialogs from Exercise 1 on.")

# 5 — UI walkthrough (with real screenshot)
s = white_slide("Cowork UI walkthrough", "Four places in the left navigation you\u2019ll use all day.")
pic = s.shapes.add_picture(UI_SHOT, Inches(0.55), Inches(1.6), width=Inches(6.4))
pic.line.color.rgb = RGBColor.from_string(W_LINE)
pic.line.width = Pt(0.75)
text(s, 0.55, 1.6 + pic.height / 914400 + 0.1, 6.4, 0.6,
     [{"runs": rich("**New task** home: Start a task box, model picker (e.g., Opus 5.5), reasoning effort "
                    "(e.g., High), attach (+), dictate, and \u201cTry these next\u201d starters.", 11, W_SUB)}])
band(s, 0.55, 5.35, 6.4, 1.35, "While a task runs, open the **side panel**: progress, skill chips, the "
     "**Output folder** of files Cowork created, and schedules update in real time.", size=13)
for i, (h, sub, acc) in enumerate([
        ("NEW TASK", "Start a task, pick a model and reasoning effort, attach files, or try a starter.", BLUE),
        ("MY TASKS", "Find and resume previous tasks \u2014 past work is always a click away.", PURPLE),
        ("AUTOMATIONS", "Run tasks on a schedule or on events. Tabs: Runs and Manage schedules.", GREEN),
        ("CUSTOMIZE", "Custom instructions, your personal skills, and plugins.", RED)]):
    y = 1.6 + i * 1.3
    box(s, 7.2, y, 5.58, 1.2, "FFFFFF", line=W_LINE, shadow=True)
    box(s, 7.2, y, 0.06, 1.2, acc)
    text(s, 7.42, y + 0.12, 5.2, 1.0, [{"runs": [(h, {"size": 14, "bold": True, "color": acc})], "space_after": 3},
                                       {"runs": rich(sub, 12.5, W_BODY)}])
notes(s, "COWORK UI WALKTHROUGH (part of 0:20-0:40). The screenshot is the NEW TASK home ('What can I do for "
         "you?'): Start a task box, model picker, reasoning effort, attach, dictate, 'Try these next'. Then "
         "MY TASKS (resume past work), AUTOMATIONS (Runs / Manage schedules), CUSTOMIZE (instructions, "
         "skills, plugins). Point out the session side panel — you'll refer to it all day. Detail: "
         "reference/07-cowork-ui-walkthrough.md.")

# 5b — How approvals work
s = white_slide("How approvals work \u2014 one at a time",
                "Cowork asks before sensitive actions: sending, posting, deleting, creating meetings.")
APPR = [("Send / Post / Create", "Approves this one action", "Yes \u2014 after reading the preview", GREEN),
        ("Cancel", "Skips it; Cowork continues with the rest", "Yes \u2014 whenever unsure", GREEN),
        ("Show parameters", "Shows recipients, targets, and details", "Yes \u2014 check before approving", GREEN),
        ("Approve All (n)", "Approves every pending action at once", "No \u2014 not today", RED),
        ("More options \u2192 Always allow", "Stops asking for similar actions this session", "No \u2014 not today", RED)]
text(s, 0.75, 1.6, 3.6, 0.4, [{"runs": [("OPTION", {"size": 12, "bold": True, "color": W_SUB})]}])
text(s, 4.55, 1.6, 4.5, 0.4, [{"runs": [("WHAT IT DOES", {"size": 12, "bold": True, "color": W_SUB})]}])
text(s, 9.25, 1.6, 3.4, 0.4, [{"runs": [("USE IT TODAY?", {"size": 12, "bold": True, "color": W_SUB})]}])
for i, (o, what, use, acc) in enumerate(APPR):
    y = 2.0 + i * 0.66
    box(s, 0.55, y, 12.23, 0.56, "FFFFFF", line=W_LINE, shadow=True)
    box(s, 0.55, y, 0.06, 0.56, acc)
    text(s, 0.75, y, 3.7, 0.56, [{"runs": [(o, {"size": 13.5, "bold": True, "color": W_TITLE})]}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 4.55, y, 4.6, 0.56, [{"runs": [(what, {"size": 13, "color": W_BODY})]}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.25, y, 3.4, 0.56, [{"runs": [(use, {"size": 13, "bold": True, "color": acc})]}], anchor=MSO_ANCHOR.MIDDLE)
band(s, 0.55, 5.45, 12.23, 1.3,
     "**Why:** one click on **Approve All** could bulk-send mail or change many items at once. **Clicked one by "
     "mistake?** Open the side panel\u2019s **Permissions** section and revoke it.",
     fill=PLUM, color="FFFFFF", size=15)
notes(s, "HOW APPROVALS WORK (end of the UI walkthrough, before Exercise 1). Show a real approval dialog on screen. "
         "Name each option: the action button, Cancel, Show parameters, and the two to avoid today: Approve All (n) "
         "and More options -> Always allow / Approve & don't ask again. Say it plainly: 'One at a time. No Approve "
         "All.' Show where to revoke: side panel -> Permissions. Source: Microsoft Learn, 'Use Copilot Cowork -> "
         "Approve actions'. Detail: reference/07-cowork-ui-walkthrough.md.")

# 5c — Workshop kit: download links
kit_links_slide()

# 6 — Setup
s = white_slide("Setup \u2014 sign in and copy the data", "Do this before Exercise 1. Proctors: triage sign-in issues now.")
wcard(s, 0.55, 1.6, 3.95, 3.3, BLUE, "1 \u00b7 SIGN IN",
      bullets(["Use the workshop account you were given", "Open copilot.cloud.microsoft in Microsoft Edge",
               "Select **Cowork** next to Chat", "Smoke test: \u201cGive me a one-sentence hello.\u201d", "Edge profile = the **workshop account** (for Ex 3)"], size=14))
wcard(s, 4.69, 1.6, 3.95, 3.3, PURPLE, "2 \u00b7 COPY THE DATA",
      bullets(["Download and extract **zava-sample-knowledge.zip**",
               "In **your** OneDrive \u2192 **Documents**, use **Folder upload** to add the **ai_hr_cowork_workshop** folder",
               "Attach files with **+ \u2192 Attach cloud files**, or type **/**"], size=14))
wcard(s, 8.83, 1.6, 3.95, 3.3, GREEN, "3 \u00b7 ETIQUETTE",
      bullets(["Your OneDrive, drafts, and skills are your own", "Keep custom skills **\u201cOnly you\u201d** or initialed",
               "Fictional Zava data only \u2014 no real PII", "Never send to real people; save drafts"], size=14))
band(s, 0.55, 5.2, 12.23, 1.55,
     "**Attendees share one tenant**, so your screens match each other and org search is consistent. The "
     "**facilitator demos from a separate tenant**, so the instructor\u2019s screen may look a little different.",
     fill=PLUM, color="FFFFFF", size=14)
notes(s, "SETUP (part of 0:20-0:40). Whole room: sign in + smoke test, then download zava-sample-knowledge.zip, "
         "extract it, and use Folder upload to add the ai_hr_cowork_workshop folder (the six Word and Excel files) "
         "to Documents in their own OneDrive; no Folder upload option? create the folder and upload the files. Show how to attach "
         "a file: + -> Attach cloud files, or type /. EDGE PROFILE CHECK: the browser task in Ex 3 runs in the Edge profile signed in with the workshop account; anyone in their own corporate profile adds a new Edge profile now (not InPrivate). Attendees share ONE tenant; YOU are on a "
         "SEPARATE tenant. Proctors triage; anyone blocked follows your demo. readiness-checklist.md has "
         "the triage, and seed-content.md has the seed data for Exercises 1 and 2.")

# 6b — Setup: custom instructions
CI_SHOT = os.path.join(ROOT, "reference", "media", "customize-instructions.png")
s = white_slide("Setup \u2014 Customize instructions for Cowork",
                "Guidance Cowork automatically adds to the start of every task. Set it once; it applies all day.")
pic = s.shapes.add_picture(CI_SHOT, Inches(0.55), Inches(1.55), width=Inches(12.23))
pic.name = "Picture CustomInstructions"
pic._element.nvPicPr.cNvPr.set("descr", "Screenshot of the Cowork card 'Customize instructions for Cowork: Give Cowork "
                               "guidance that is automatically added to the start of every task.'")
wcard(s, 0.55, 3.5, 4.3, 2.4, BLUE, "HOW",
      bullets(["**Customize \u2192 Preferences**", "Select **Customize instructions for Cowork**",
               "Paste the text, then **save**", "Test it in a **new task**"], size=14, kind="num"))
wcard(s, 5.05, 3.5, 7.73, 2.4, PURPLE, "PASTE THIS (WORKBOOK SETUP STEP D)",
      [{"runs": [("I work in HR at Zava. Write in a warm, professional, inclusive tone suitable for employee "
                  "communications. When you answer a policy or benefits question, cite the source document and add "
                  "\u201cPolicies can change \u2014 please confirm with HR.\u201d Save emails and messages as drafts for "
                  "me to review; during this workshop, never send anything to anyone but me. Use only the Zava sample "
                  "files in my OneDrive folder Documents/ai_hr_cowork_workshop, and never include real employee personal data.",
                  {"size": 13.5, "italic": True, "color": W_BODY})]}])
band(s, 0.55, 6.05, 12.23, 0.75,
     "**Personal to you** \u00b7 type **/** to reference a file or person \u00b7 up to ~20 KB, but **shorter is better** "
     "\u00b7 Role samples: workbook step D", fill=PLUM, color="FFFFFF", size=13)
notes(s, "CUSTOM INSTRUCTIONS (part of 0:20-0:40, about 3 min). Show it live: Customize -> Preferences -> 'Customize "
         "instructions for Cowork'. Explain: guidance Cowork automatically adds to the START of every task, so tone, "
         "format, and rules don't have to be repeated in each prompt. Everyone pastes the workbook text (Setup step D), "
         "saves, then tests in a new task: 'Draft a two-sentence reminder to employees that open enrollment is in "
         "November.' The reply should use the tone and end with the confirm-with-HR note. Instructions are personal "
         "(neighbors in the shared tenant don't see them), support rich text and / references, up to about 20 KB; "
         "keep them short because they're included in every task. Source: Microsoft Learn, 'Customize Copilot Cowork "
         "-> Custom instructions in Cowork'. MORE SAMPLES: workbook step D has role-based samples to paste back at work "
         "(HR business partner, recruiter, HR operations, employee comms); the full set is in "
         "reference/02-settings-and-models.md. Tell people to pick ONE, not stack them.")

# 7 — Settings & models
s = white_slide("Settings & models", "Shape how Cowork works \u2014 you won\u2019t change these every task.")
tiles = [("MODEL PICKER", "Pick a model or let Cowork decide (default).", BLUE),
         ("REASONING EFFORT", "Balance quality, speed, and cost per task.", PURPLE),
         ("CUSTOM INSTRUCTIONS", "Preferences Cowork applies to every task.", GREEN),
         ("PLUGINS", "Extend Cowork from the Microsoft 365 App Store.", RED),
         ("AUTOMATIONS", "Scheduled prompts and event-driven tasks.", OLIVE),
         ("IMAGE GENERATION", "Simple visuals for comms and slides.", BLUE)]
for i, (h, sub, acc) in enumerate(tiles):
    wcard(s, 0.55 + (i % 3) * 4.14, 1.6 + (i // 3) * 1.85, 3.95, 1.7, acc, h, [{"runs": rich(sub, 14, W_BODY)}], hsize=14)
band(s, 0.55, 5.45, 12.23, 1.3,
     "**Tip:** for sensitive employee writing, use a higher-capability model and higher reasoning effort. "
     "For routine drafts, the defaults are fine. Cowork is usage-billed \u2014 match effort to the task.")
notes(s, "SETTINGS & MODELS. Model picker ('let Cowork decide' is a fine default; Anthropic models may appear "
         "as a subprocessor depending on tenant) and reasoning effort (quality vs speed vs cost). Set one "
         "custom instruction as a class in Customize. Detail: reference/02-settings-and-models.md.")

# 8 — Skills
s = white_slide("Skills \u2014 how Cowork gets work done",
                "Cowork loads skills automatically and shows them as chips in the side panel.")
SKILLS = ["Word", "Excel", "PowerPoint", "PDF", "Email", "Scheduling", "Calendar", "Meetings", "Daily Briefing",
          "Enterprise Search", "Deep Research", "Communications", "Adaptive Cards", "App (Frontier)"]
USED = {"Word", "Excel", "PowerPoint", "Email", "Scheduling", "Calendar", "Deep Research", "Communications",
        "Daily Briefing"}
for i, name in enumerate(SKILLS):
    x, y = 0.55 + (i % 5) * 2.47, 1.6 + (i // 5) * 0.85
    acc = PURPLE if name in USED else GRAY
    box(s, x, y, 2.3, 0.7, "FFFFFF", line=W_LINE, shadow=True)
    box(s, x, y, 0.06, 0.7, acc)
    text(s, x + 0.2, y, 2.0, 0.7, [{"runs": [(name, {"size": 13.5, "bold": True, "color": W_TITLE})]}],
         anchor=MSO_ANCHOR.MIDDLE)
band(s, 0.55, 4.35, 12.23, 1.15,
     "**Three kinds of skills:** built-in (above) \u00b7 **custom** (you build one in Exercise 4) \u00b7 **plugins** "
     "from the Microsoft 365 App Store. Purple = used hands-on today.", size=14)
band(s, 0.55, 5.7, 12.23, 1.05,
     "App (Frontier) requires the Frontier program \u2014 awareness only in this workshop.", fill="F5F5F5",
     color=W_SUB, size=13)
notes(s, "SKILLS. You don't invoke skills manually — Cowork activates them and shows chips. Ex 1 gathers mail, "
         "calendar, and chats through Work IQ; Ex 2 Deep Research + Word/Excel; Ex 3 is browser use (no skill chip); Ex 5 Word/Excel; Ex 6 PowerPoint, Scheduling, Communications; Ex 7 Daily Briefing. "
         "Full list: reference/03-skills-catalog.md.")

# 9 — Responsible use
s = white_slide("Responsible use \u2014 the golden rules", "HR handles sensitive people data. Make these non-negotiable.")
RULES = [("Every output is a draft", "Review before you send, share, or file \u2014 especially employee comms.", BLUE),
         ("Approve at checkpoints", "Cowork pauses before irreversible actions. Read before you confirm.", PURPLE),
         ("Use fictional data today", "All sample files are the fictional Zava. No real PII.", GREEN),
         ("Cite & confirm", "For policy answers, name the source and note HR should confirm.", RED),
         ("Right person, right data", "Cowork only reaches content you already have permission to see.", OLIVE)]
for i, (h, sub, acc) in enumerate(RULES):
    y = 1.6 + i * 1.03
    box(s, 0.55, y, 12.23, 0.92, "FFFFFF", line=W_LINE, shadow=True)
    box(s, 0.55, y, 0.06, 0.92, acc)
    text(s, 0.8, y, 0.5, 0.92, [{"runs": [(str(i + 1), {"size": 22, "bold": True, "color": acc})]}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.35, y, 3.9, 0.92, [{"runs": [(h, {"size": 16, "bold": True, "color": W_TITLE})]}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 5.3, y, 7.3, 0.92, [{"runs": rich(sub, 14, W_SUB)}], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "RESPONSIBLE USE. Draft -> review -> approve. Fictional Zava data only; no real PII; never send to real "
         "people (send only to yourself). Cite sources + 'confirm with HR'. Exercise 1's prompt also "
         "keeps recommendations on workstreams, not on evaluating individuals. Detail: "
         "reference/06-responsible-use.md.")

# 9b — Prompting best practices: Goal · Source · Expectations · Constraints (same colors as the Word workbook)
s = white_slide("Prompting best practices",
                "Strong prompts include four elements. No labels needed: just make sure each one is there.")
for i, key in enumerate("GSEC"):
    name, fill, dark, ask = PROMPT_EL[key]
    x = 0.55 + i * 3.1
    box(s, x, 1.55, 2.95, 1.15, fill)
    box(s, x, 1.55, 0.07, 1.15, dark)
    text(s, x + 0.25, 1.62, 2.6, 0.4, [{"runs": [(name, {"size": 17, "bold": True, "color": dark})]}])
    text(s, x + 0.25, 2.03, 2.6, 0.65, [{"runs": [(ask, {"size": 12, "color": W_BODY})]}], line_spacing=1.0)
# Weak
box(s, 0.55, 2.9, 12.23, 0.95, "FFFFFF", line=W_LINE, shadow=True)
box(s, 0.55, 2.9, 0.07, 0.95, RED)
text(s, 0.85, 2.9, 1.2, 0.95, [{"runs": [("WEAK", {"size": 15, "bold": True, "color": RED})]}], anchor=MSO_ANCHOR.MIDDLE)
text(s, 2.05, 2.9, 4.6, 0.95, [{"runs": [("“Write an email about open enrollment.”",
                                           {"size": 16, "italic": True, "color": W_TITLE})]}], anchor=MSO_ANCHOR.MIDDLE)
text(s, 6.85, 2.9, 5.75, 0.95, [{"runs": rich("Cowork has to **guess**: who it’s for, which dates apply, how long it "
                                              "should be, and whether to send it.", 13, W_SUB)}],
     anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
# Strong (same topic), each element shaded in its color and tagged for readers who can't rely on color
box(s, 0.55, 4.0, 12.23, 2.8, "FFFFFF", line=W_LINE, shadow=True)
box(s, 0.55, 4.0, 0.07, 2.8, GREEN)
text(s, 0.85, 4.18, 1.2, 0.4, [{"runs": [("STRONG", {"size": 15, "bold": True, "color": GREEN})]}])
STRONG = [("G", "Draft a reminder email to all Zava employees so that anyone who wants to change their benefits "
                "does it during November open enrollment."),
          ("S", "Use the Enrollment Windows section of benefits-summary.docx and the Who to Contact section of "
                "employee-handbook-excerpt.docx."),
          ("E", "Keep it under 150 words, with a subject line of eight words or fewer, three short bullets on what "
                "to do, and who to contact for help. Use a warm, plain-language tone."),
          ("C", "Don’t invent dates, deadlines, or plan details that aren’t in the files. Save it as a draft "
                "for me to review; don’t send it.")]
runs, keys = [], []
for i, (key, seg) in enumerate(STRONG):
    name, fill, dark, _ = PROMPT_EL[key]
    runs += [(name.upper() + " ", {"size": 9.5, "bold": True, "color": dark}), (seg, {"size": 15, "color": W_TITLE})]
    keys += [key, key]
    if i < len(STRONG) - 1:
        runs.append((" ", {"size": 15}))
        keys.append(None)
tb = text(s, 2.05, 4.18, 10.5, 2.5, [{"runs": runs}], line_spacing=1.12)
for r, key in zip(tb.text_frame.paragraphs[0].runs, keys):
    if key:
        highlight(r._r, PROMPT_EL[key][1])
text(s, 2.05, 6.22, 10.5, 0.4, [{"runs": rich("**Same topic.** Now Cowork knows what to do, where to look, "
                                              "what good looks like, and what not to do.", 13, W_SUB)}])
notes(s, "PROMPTING BEST PRACTICES (about 4 min, before the exercises). Read the WEAK prompt aloud and ask the room: "
         "what would Cowork have to guess? (Audience, which dates, length, tone, whether to send.) Then walk the STRONG "
         "version element by element: GOAL (blue) says what outcome and for whom; SOURCE (green) names the exact files "
         "and sections; EXPECTATIONS (orange) describe what good looks like: length, subject line, bullets, tone; "
         "CONSTRAINTS (purple) say what not to do. Point out 'don't invent dates': the benefits summary only says "
         "November, so a weak prompt invites a made-up deadline. Key message: no labels and no fixed order; just check "
         "all four are there. Context about the situation belongs in the Goal. In the Word workbook every exercise "
         "prompt from Ex 2 onward is color-coded with these same colors (setup step F has the legend and this example); "
         "Exercise 1's prompt is left as is. When Cowork asks a clarifying question or misses, the missing piece is "
         "usually one of the four: add it in a follow-up.")

# 10 — How to read an exercise card (the template itself, annotated)
scenario_card(GUIDE_CARD)

def output_folder_slide():
    s = clone_slide(S_CARD)
    keep_only(s, {"Title 3", "NavPill_SitesPages", "TextBox 29", "TextBox 30", "TextBox 49"})
    d = by_name(s)
    set_first_run(d["Title 3"], "Find your outputs \u00b7 the Output folder")
    set_first_run(d["NavPill_SitesPages"], "Ex 01")
    set_first_run(d["TextBox 30"], "Time")
    set_first_run(d["TextBox 49"], "3 min")
    # Steps panel (left)
    box(s, 0.46, 1.3, 7.3, 5.75, PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    pill(s, 4.11, 1.15, "Show the room")
    steps = ["**Open the side panel** \u2014 select the side panel toggle at the right of the session.",
             "**Input folder** = files you attached. **Output folder** = files Cowork created.",
             "Select **Preview** on the HTML dashboard: it opens in a **split view**. Try the "
             "**full-screen toggle** and **Open in native app**.",
             "**Download** saves one file; **Download All** saves every output as a zip.",
             "Open **OneDrive \u2192 Cowork** in a new tab: every output is saved there too.",
             "Now **everyone** does it before we move on."]
    text(s, 0.85, 1.72, 6.6, 4.1, bullets(steps, size=14, color=INK, kind="num"), font=SEG_DISP,
         space_after=8, line_spacing=1.05)
    box(s, 0.75, 5.95, 6.72, 0.9, PLUM, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(s, 1.0, 5.95, 6.3, 0.9, [{"runs": [("Every exercise that creates a file ends here. ",
                                            {"size": 12.5, "bold": True, "color": "FFFFFF", "font": SEG_SEMI}),
                                           ("Can\u2019t find it? Side panel toggle \u2192 Output folder, or "
                                            "OneDrive \u2192 Cowork.", {"size": 12.5, "color": "FFFFFF", "font": SEG_DISP})]}],
         anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.03)
    # Side panel mock-up (right), sections as documented on Microsoft Learn
    x0, y0, w = 8.2, 1.3, 4.9
    box(s, x0, y0, w, 5.75, "FFFFFF", line=W_LINE, shadow=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
    text(s, x0 + 0.25, y0 + 0.12, w - 0.5, 0.35, [{"runs": [("Side panel", {"size": 12, "bold": True, "color": W_SUB})]}])
    rows = [("Progress", "Progress bar and step-by-step log", False),
            ("Input folder", "Files you attached", False),
            ("Output folder", None, True),
            ("Skills", "Chips for the skills Cowork loaded", False),
            ("Schedule", "Weekday run \u2014 edit, pause, resume, delete", False),
            ("Permissions", "\u201cDon\u2019t ask again\u201d approvals \u2014 revoke here", False)]
    y = y0 + 0.55
    for name, val, hi in rows:
        h = 1.15 if hi else 0.66
        if hi:
            box(s, x0 + 0.15, y, w - 0.3, h, "FBF6FD", line=PURPLE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
            text(s, x0 + w - 1.55, y + 0.08, 1.2, 0.28,
                 [{"runs": [("Download All", {"size": 9, "bold": True, "color": BLUE})], "align": PP_ALIGN.RIGHT}])
        text(s, x0 + 0.3, y + 0.06, w - 0.6, 0.3,
             [{"runs": [(name, {"size": 11.5, "bold": True, "color": PURPLE if hi else W_TITLE})]}])
        if hi:
            fy = y + 0.46
            box(s, x0 + 0.3, fy, w - 0.6, 0.52, "FFFFFF", line=W_LINE)
            text(s, x0 + 0.42, fy, 2.5, 0.52, [{"runs": [("executive-command-center.html", {"size": 9.5, "color": W_BODY})]}],
                 anchor=MSO_ANCHOR.MIDDLE)
            for k, lbl in enumerate(("Preview", "Download")):
                bx = x0 + w - 1.9 + k * 0.82
                box(s, bx, fy + 0.11, 0.76, 0.3, PILL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.4)
                text(s, bx, fy + 0.11, 0.76, 0.3, [{"runs": [(lbl, {"size": 8.5, "bold": True, "color": INK})],
                                                    "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
        else:
            text(s, x0 + 0.3, y + 0.34, w - 0.6, 0.28, [{"runs": [(val, {"size": 9.5, "color": W_SUB})]}])
        y += h + 0.1
    text(s, x0 + 0.25, y0 + 5.4, w - 0.5, 0.3,
         [{"runs": [("Illustration of the side panel\u2019s sections, not a screenshot",
                     {"size": 8.5, "italic": True, "color": W_SUB})]}])
    notes(s, "FIND YOUR OUTPUTS (Ex 1, ~3 min). Take your time here; every later exercise relies on it. On "
             "your screen: open the side panel with the side panel toggle; point out Input folder (what you "
             "attached) vs Output folder (what Cowork created, each file with Preview and Download); Preview "
             "the HTML dashboard (split view, full-screen toggle, Open in native app); show Download and "
             "Download All (zip); open OneDrive -> Cowork in a new tab to show the same file. Then have EVERY "
             "attendee do it before moving on; proctors help anyone who can't find the toggle. The panel on "
             "the right is an illustration of the side panel's sections (per Microsoft Learn), not a screenshot.")
    return s


PLUGINS_SHOT = os.path.join(ROOT, "reference", "media", "customize-plugins.png")


def plugins_spotlight_slide():
    s = clone_slide(S_CARD)
    keep_only(s, {"Title 3", "NavPill_SitesPages", "TextBox 29", "TextBox 30", "TextBox 49"})
    d = by_name(s)
    set_first_run(d["Title 3"], "Spotlight \u00b7 Extend Cowork with HR plugins")
    set_first_run(d["NavPill_SitesPages"], "Talk only")
    set_first_run(d["TextBox 30"], "Time")
    set_first_run(d["TextBox 49"], "5 min")
    # Screenshot (left)
    box(s, 0.46, 1.3, 7.3, 5.75, PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    pill(s, 4.11, 1.15, "Customize \u2192 Plugins")
    pic = s.shapes.add_picture(PLUGINS_SHOT, Inches(0.7), Inches(1.6), width=Inches(6.82))
    pic.name = "Picture Plugins"
    pic._element.nvPicPr.cNvPr.set(
        "descr", "Screenshot from Microsoft Learn of Copilot Cowork Customize, Plugins tab, listing installed "
                 "plugins with on/off toggles, plus Discover and Upload plugin options.")
    text(s, 0.7, 1.6 + 6.82 * 541 / 1030 + 0.05, 6.82, 0.3,
         [{"runs": [("Microsoft Learn screenshot (Dynamics 365 plugins shown). HR plugins appear in the same list.",
                     {"size": 9, "italic": True, "color": W_SUB})]}])
    lk = text(s, 0.7, 1.6 + 6.82 * 541 / 1030 + 0.3, 6.82, 0.28,
              [{"runs": [("Full catalog: learn.microsoft.com/microsoft-365/copilot/cowork/cowork-available-plugins",
                          {"size": 9, "color": BLUE})]}])
    lk.text_frame.paragraphs[0].runs[0].hyperlink.address = (
        "https://learn.microsoft.com/en-us/microsoft-365/copilot/cowork/cowork-available-plugins")
    box(s, 0.75, 5.8, 6.72, 0.95, PLUM, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(s, 1.0, 5.8, 6.3, 0.95, [{"runs": [("A plugin = skills + connectors. ",
                                            {"size": 12.5, "bold": True, "color": "FFFFFF", "font": SEG_SEMI}),
                                           ("Connectors can be MCP servers, so Cowork can read and act in "
                                            "your HR systems \u2014 still with your approval.",
                                            {"size": 12.5, "color": "FFFFFF", "font": SEG_DISP})]}],
         anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.03)
    # HR plugins + benefits (right)
    x0, y0, w = 8.2, 1.3, 4.9
    box(s, x0, y0, w, 5.75, "FFFFFF", line=W_LINE, shadow=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
    text(s, x0 + 0.25, y0 + 0.12, w - 0.5, 0.35,
         [{"runs": [("HR plugin examples", {"size": 12, "bold": True, "color": W_SUB})]}])
    rows = [("Gusto", "Payroll, PTO and employee records"),
            ("ZipRecruiter \u00b7 Dice.com", "Job posts and candidate pipelines"),
            ("Cronofy MCP", "Interview scheduling across calendars"),
            ("Articulate", "Learning and training content"),
            ("\u2b50 SAP SuccessFactors", "No catalog plugin yet: ESS agent, or a custom MCP plugin")]
    y = y0 + 0.5
    for name, val in rows:
        hi = "SuccessFactors" in name
        if hi:
            box(s, x0 + 0.15, y - 0.06, w - 0.3, 0.64, "FBF6FD", line=PURPLE,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        text(s, x0 + 0.3, y, w - 0.6, 0.28, [{"runs": [(name, {"size": 11.5, "bold": True,
                                                              "color": PURPLE if hi else W_TITLE})]}])
        text(s, x0 + 0.3, y + 0.27, w - 0.6, 0.26, [{"runs": [(val, {"size": 9.5, "color": W_SUB})]}])
        y += 0.68
    y += 0.05
    box(s, x0 + 0.15, y, w - 0.3, 5.75 + y0 - y - 0.15, "FBF6FD", line=PURPLE,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    ben = ["**Fewer tabs** \u2014 HR data flows into the task",
           "**Governed** \u2014 admins deploy; activity in Purview audit",
           "**Your sign-in** \u2014 connectors use your own permissions",
           "**Per task** \u2014 toggle in Sources & Skills"]
    text(s, x0 + 0.3, y + 0.08, w - 0.6, 0.3, [{"runs": [("Why it matters", {"size": 11.5, "bold": True, "color": PURPLE})]}])
    text(s, x0 + 0.3, y + 0.4, w - 0.55, 1.6, bullets(ben, size=10.5, color=INK), space_after=3, line_spacing=1.0)
    notes(s, "SPOTLIGHT \u2014 HR PLUGINS (3:30-3:35). TALK ONLY, no hands-on. Show Customize \u2192 Plugins: "
             "Installed (toggles), a plugin's detail page (its skills and connectors), Discover, and Upload plugin. "
             "A plugin bundles skills and connectors; connectors can be MCP servers, so Cowork can read and act in "
             "systems like payroll, ATS or your HRIS. Example (illustrative): with a payroll plugin such as Gusto, "
             "ticket T-2008 (Owen Wright's payroll question) could be answered from the real pay record, with Cowork "
             "asking approval before any change. Each user signs in to a connector once; admins can deploy plugins "
             "('Managed by your organization') and activity is in Purview audit logs. Do NOT install plugins in the "
             "shared attendee tenant \u2014 demo from the facilitator tenant if you have one installed. "
             "SAP SUCCESSFACTORS: no SuccessFactors plugin in the Cowork catalog yet (checked Oct 1, 2026). Point to "
             "the Employee Self-Service agent with its SAP SuccessFactors extension pack (a Copilot agent, not a "
             "Cowork plugin), or a custom plugin with an MCP connector built by IT. Re-check the catalog before each "
             "delivery: learn.microsoft.com/microsoft-365/copilot/cowork/cowork-available-plugins. "
             "Details: reference/09-plugins.md.")
    return s


# Exercises 1-8 with breaks after Ex 2 and Ex 5
NEXT_AFTER = {2: "Ex 3 \u00b7 Navigate websites", 5: "Ex 6 \u00b7 Onboarding pack"}
for ex in EXERCISES:
    mark_section(f"Ex {ex['num']} \u00b7 " + {1: "Executive Command Center", 2: "Deep Research", 3: "Navigate websites",
                                           4: "Custom skill", 5: "Recruiting + reporting", 6: "Onboarding pack",
                                           7: "Automate & share", 8: "HR Policy Agent"}.get(ex["num"], ex["short"]))
    exercise_divider(ex)
    scenario_card(ex)
    hands_on(ex)
    if ex["num"] == 1:
        output_folder_slide()
    if ex["num"] == 7:
        mark_section("Spotlight \u00b7 HR plugins")
        plugins_spotlight_slide()
    if ex["num"] in NEXT_AFTER:
        b = break_slide(NEXT_AFTER[ex["num"]])
        notes(b, "BREAK (10 min). Sweep the room for blockers. " +
              ("Next: Cowork drives a web browser, then the centerpiece custom-skill build. Proctors: check Edge profiles and the Edge Cowork setting now." if ex["num"] == 2 else
               "Check everyone's Exercise 4 skill saved — Exercise 7 shares it."))

# Facilitation
mark_section("Close")
s = white_slide("Facilitation & troubleshooting", "For instructors \u2014 skip live or use during a break.")
wcard(s, 0.55, 1.6, 5.98, 3.35, BLUE, "KEEP 25 PEOPLE TOGETHER",
      bullets(["1 facilitator + 1\u20132 floaters", "Use each hands-on slide\u2019s **checkpoint** to sync",
               "Load **seed-content.md** within 24 h of the session", "Check results with the **answer key**", "Watch the clock: **time checks** at each break"], size=14))
wcard(s, 6.8, 1.6, 5.98, 3.35, PURPLE, "COMMON BLOCKERS",
      bullets(["No Cowork toggle \u2192 spare licensed account", "Acted without asking \u2192 revoke in side panel **Permissions**",
               "Empty dashboard \u2192 expected; show your demo", "No browser task \u2192 Edge, signed in as the workshop account", "Agent ignores files \u2192 wait for \u201cPreparing\u201d, then refresh"], size=14))
band(s, 0.55, 5.25, 12.23, 1.5,
     "**Fallback:** anyone blocked follows your demo, using **facilitator-answer-key.md**. Whole room down? Stop after "
     "10 minutes and switch to **Plan B** (readiness-checklist.md). Running long? Use the time checks in the facilitator guide.",
     fill=PLUM, color="FFFFFF", size=14)
notes(s, "FACILITATION. Biggest risks with ~25 people: pace variance, account readiness, and sparse data in new "
         "accounts (Ex 1, fixed by seed-content.md). Attendees share one tenant; you're on a separate one. "
         "Mirrors facilitator-guide.md, facilitator-answer-key.md, and readiness-checklist.md (timeline, cost "
         "planning, Plan B). Time checks: 0:40 Ex 1, 1:25 break, 2:00 Ex 4, 2:45 break, 3:15 Ex 7, 3:35 Ex 8; cut in that order (shorter breaks, demo the Ex 3 scorecard, 6a + 6c only, 7a only, Ex 8 as a demo). Never cut Ex 4.")

# Knowledge check
s = white_slide("Quick knowledge check", "Show of hands \u2014 answers in after-the-workshop.md.")
KC = [("1", "You need a 3-bullet summary of one policy. Which tool?", "Copilot Chat", BLUE),
      ("2", "Cowork shows **Approve All (5)**. What happens if you click it?", "All five run, with no individual review", PURPLE),
      ("3", "Why can\u2019t you add **.md** files as Agent Builder knowledge?", "Unsupported \u2014 use .docx, .pdf, or .xlsx", GREEN),
      ("4", "Skill vs. agent \u2014 who benefits?", "A skill helps **you**; an agent helps **others**", RED)]
for i, (n, q, a, acc) in enumerate(KC):
    y = 1.6 + i * 1.05
    box(s, 0.55, y, 12.23, 0.92, "FFFFFF", line=W_LINE, shadow=True)
    box(s, 0.55, y, 0.06, 0.92, acc)
    text(s, 0.8, y, 0.5, 0.92, [{"runs": [(n, {"size": 24, "bold": True, "color": acc})]}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.4, y, 6.9, 0.92, [{"runs": rich(q, 15, W_TITLE)}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.5, y, 4.1, 0.92, [{"runs": rich(a, 13, W_SUB)}], anchor=MSO_ANCHOR.MIDDLE)
band(s, 0.55, 5.9, 12.23, 0.85,
     "**Reveal answers one at a time** \u2014 the right column is for you. Tomorrow: the full 8-question check, survey, and 30-day plan.",
     size=13)
notes(s, "KNOWLEDGE CHECK (start of wrap-up). Read each question, take a show of hands, THEN reveal the answer (the "
         "right column; consider covering it or using an animation). Then return to the six learning objectives. "
         "The full 8-question check, feedback survey, and 30-day adoption plan are in after-the-workshop.md; send "
         "it the next day.")

# Wrap-up
s = white_slide("Wrap-up & next steps", "Three tools, eight exercises, one habit: draft \u2192 review \u2192 approve.")
TAKE = [("Three tools", "Copilot Chat for quick answers, **Cowork** to get multi-step work done, **Agent Builder** for a reusable helper.", BLUE),
        ("What you built", "A command center, a cited briefing, a browser task across two sites, a custom skill, an onboarding pack, an automation, and an agent.", PURPLE),
        ("Take it with you", "The **quick-reference card**, the prompt library, and the scenario cards in your workbook.", GREEN),
        ("Next week", "Pick **one** task to hand to Cowork. Tomorrow you\u2019ll get a 30-day adoption plan.", RED)]
for i, (h, sub, acc) in enumerate(TAKE):
    y = 1.6 + i * 1.0
    box(s, 0.55, y, 12.23, 0.88, "FFFFFF", line=W_LINE, shadow=True)
    box(s, 0.55, y, 0.06, 0.88, acc)
    text(s, 0.8, y, 2.6, 0.88, [{"runs": [(h, {"size": 15, "bold": True, "color": acc})]}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.45, y, 9.2, 0.88, [{"runs": rich(sub, 14, W_BODY)}], anchor=MSO_ANCHOR.MIDDLE)
band(s, 0.55, 5.75, 12.23, 1.0,
     "**Clean up:** pause or delete the schedules you created in Exercises 1 and 7 (Automations \u2192 Manage schedules).",
     fill=PLUM, color="FFFFFF", size=14)
notes(s, "WRAP-UP (3:55-4:00). Recap the THREE tools and everything they built. Reinforce the golden rules. Ask "
         "each person to name ONE task for next week. Remind them to pause/delete the Ex 1 and Ex 7 schedules.")

# Thank you
s = title_slide("Thank you", "Questions? Try the stretch prompts, then explore the kit\u2019s README.")
notes(s, "CLOSE. Thank the group and take questions. Point to README.md, the prompt library, and the "
         "custom-skill and Agent Builder guides.")

# ---------------------------------------------------------------- remove template slides, renumber, sections
sld_lst = prs.slides._sldIdLst
for sid in list(sld_lst)[:N_ORIG]:
    rId = sid.rId
    sld_lst.remove(sid)
    prs.part.drop_rel(rId)
ids = []
for i, sid in enumerate(sld_lst):
    sid.set("id", str(256 + i))
    ids.append(str(256 + i))

total = len(ids)
SECTIONS = [(name, start, (SECTION_MARKS[i + 1][1] - 1) if i + 1 < len(SECTION_MARKS) else total)
            for i, (name, start) in enumerate(SECTION_MARKS)]
for sec_lst in prs.part._element.iter("{%s}sectionLst" % P14):
    for c in list(sec_lst):
        sec_lst.remove(c)
    for name, a, b in SECTIONS:
        sec = etree.SubElement(sec_lst, "{%s}section" % P14, name=name, id="{%s}" % str(uuid.uuid4()).upper())
        lst = etree.SubElement(sec, "{%s}sldIdLst" % P14)
        for sid in ids[a - 1:b]:
            etree.SubElement(lst, "{%s}sldId" % P14, id=sid)

prs.core_properties.title = "Getting Things Done with Copilot Cowork for HR Tasks"
prs.core_properties.author = "Christophe Ragage"
prs.core_properties.last_modified_by = "Christophe Ragage"
prs.core_properties.subject = DECK_SUBJECT

# ---------------------------------------------------------------- accessibility pass
A16 = "http://schemas.microsoft.com/office/drawing/2014/main"
ADEC = "http://schemas.microsoft.com/office/drawing/2017/decorative"
DEC_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
ALT = {"UI_SHOT": "Screenshot of the Copilot Cowork New task home page: a 'What can I do for you?' heading, "
                   "a Start a task box with model picker and reasoning effort, and 'Try these next' starter cards.",
       "m365": "Microsoft 365 data icon", "web": "Web icon", "onedrive": "OneDrive icon", "sharepoint": "SharePoint icon"}


def mark_decorative(cNvPr):
    ext_lst = cNvPr.find(qn("a:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(cNvPr, qn("a:extLst"))
    if any(e.get("uri") == DEC_URI for e in ext_lst):
        return
    ext = etree.SubElement(ext_lst, qn("a:ext"), uri=DEC_URI)
    etree.SubElement(ext, "{%s}decorative" % ADEC, val="1")


def title_first(slide):
    tree = slide.shapes._spTree
    titles = [sp for sp in tree if sp.tag == qn("p:sp") and
              (sp.find(".//" + qn("p:ph")) is not None and sp.find(".//" + qn("p:ph")).get("type") in ("title", "ctrTitle")
               or (sp.find(".//" + qn("p:cNvPr")) is not None and sp.find(".//" + qn("p:cNvPr")).get("name") in ("Title 3", "TextBox 1")))]
    anchor = tree.find(qn("p:grpSpPr"))
    for t in reversed(titles):
        anchor.addnext(t)


a11y = {"alt": 0, "decorative": 0}
for s in prs.slides:
    for el in s._element.iter(qn("p:pic")):
        c = el.find(".//" + qn("p:cNvPr"))
        blip = el.find(".//" + qn("a:blip"))
        target = s.part.related_part(blip.get(qn("r:embed"))) if blip is not None else None
        pname = str(target.partname) if target is not None else ""
        if not c.get("descr"):
            desc = ALT["UI_SHOT"] if el.getparent() is s.shapes._spTree and c.get("name", "").startswith("Picture") and "image" in pname and el.find(".//" + qn("a:ext")) is not None and int(el.find(".//" + qn("a:xfrm")).find(qn("a:ext")).get("cx")) > 3000000 else None
            if desc is None:
                grp_text = " ".join(t.text or "" for t in el.getparent().iter(qn("a:t")))
                for k, label in (("m365", "M365"), ("web", "Web"), ("onedrive", "OneDrive"), ("sharepoint", "SharePoint")):
                    if label in grp_text:
                        desc = ALT[k]
                        break
            if desc:
                c.set("descr", desc); a11y["alt"] += 1
            else:
                mark_decorative(c); a11y["decorative"] += 1
    for sp in s.shapes._spTree.iter(qn("p:sp")):
        txb = sp.find(qn("p:txBody"))
        has_text = txb is not None and "".join(t.text or "" for t in txb.iter(qn("a:t"))).strip()
        is_ph = sp.find(".//" + qn("p:ph")) is not None
        if not has_text and not is_ph:
            mark_decorative(sp.find(".//" + qn("p:cNvPr"))); a11y["decorative"] += 1
    title_first(s)
print("accessibility:", a11y)
prs.save(OUT)

# The template's docProps/app.xml still lists the template's own title; replace it with this deck's.
import re, shutil, tempfile, zipfile
with zipfile.ZipFile(OUT) as zin:
    app = zin.read("docProps/app.xml").decode("utf-8")
    app = re.sub(r"<vt:lpstr>[^<]*(?:Nifty|Fifty)[^<]*</vt:lpstr>",
                 "<vt:lpstr>Getting Things Done with Copilot Cowork for HR Tasks</vt:lpstr>", app)
    tmp = tempfile.mktemp(suffix=".pptx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = app.encode("utf-8") if item.filename == "docProps/app.xml" else zin.read(item.filename)
            zout.writestr(item, data)
shutil.move(tmp, OUT)
print("saved", OUT, "slides:", total)
