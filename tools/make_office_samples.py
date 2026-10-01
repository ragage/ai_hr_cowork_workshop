"""Convert the Zava sample files to Agent Builder-friendly .docx / .xlsx copies."""
import csv
import pathlib
import re
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SRC = pathlib.Path(__file__).resolve().parent.parent / "sample-knowledge"
INK = RGBColor(0x24, 0x24, 0x24)
ACCENT = RGBColor(0x3B, 0x10, 0x41)
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|_[^_\s][^_]*_)")


def add_runs(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"
        elif part[0] in "*_" and part[-1] == part[0] and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            par.add_run(part)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def unwrap(lines):
    """Join soft-wrapped continuation lines into their list item / paragraph / quote."""
    out = []
    for ln in lines:
        s = ln.rstrip()
        starts_block = (not s.strip() or s.startswith(("#", "|", ">")) or re.match(r"^\s*(- |\d+\. )", s))
        if out and s.strip() and not starts_block and out[-1].strip() and not out[-1].startswith(("#", "|")):
            out[-1] = out[-1] + " " + s.strip()
        elif out and s.startswith(">") and out[-1].startswith(">") and s.lstrip("> ").strip():
            out[-1] = out[-1] + " " + s.lstrip("> ").strip()
        else:
            out.append(s)
    return out


def md_to_docx(md_path):
    lines = unwrap(md_path.read_text(encoding="utf-8").splitlines())
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Segoe UI"; st.font.size = Pt(10.5); st.font.color.rgb = INK
    for lvl, size in ((1, 20), (2, 14), (3, 12)):
        h = doc.styles["Heading %d" % lvl]
        h.font.name = "Segoe UI Semibold"; h.font.size = Pt(size); h.font.color.rgb = ACCENT; h.font.bold = False
    title = None
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1; continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT
            for r_i, row in enumerate(rows):
                for c_i, val in enumerate(row):
                    cell = t.cell(r_i, c_i)
                    cell.paragraphs[0].text = ""
                    add_runs(cell.paragraphs[0], val)
                    if r_i == 0:
                        shade(cell, "F3E5FD")
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,3}) (.*)", ln)
        if m:
            text = m.group(2)
            if len(m.group(1)) == 1 and title is None:
                title = text
            doc.add_heading(text, level=len(m.group(1)))
        elif ln.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            add_runs(p, ln.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True; r.font.color.rgb = RGBColor(0x61, 0x61, 0x61)
        elif re.match(r"^\s*- \[ \] ", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, "\u2610 " + re.sub(r"^\s*- \[ \] ", "", ln))
        elif re.match(r"^\s*- ", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^\s*- ", "", ln))
        else:
            p = doc.add_paragraph()
            add_runs(p, ln)
        i += 1
    doc.core_properties.title = title or md_path.stem
    doc.core_properties.author = "Zava HR (fictional sample)"
    out = md_path.with_suffix(".docx")
    doc.save(out)
    return out


def csv_to_xlsx(csv_path, sheet):
    wb = Workbook(); ws = wb.active; ws.title = sheet
    with csv_path.open(encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    for r in rows:
        ws.append([int(v) if v.isdigit() else v for v in r])
    for c in ws[1]:
        c.font = Font(bold=True)
    for idx, _ in enumerate(rows[0], 1):
        width = max(len(str(r[idx - 1])) for r in rows) + 3
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top")
    ref = "A1:%s%d" % (get_column_letter(len(rows[0])), len(rows))
    tbl = Table(displayName=sheet.replace(" ", ""), ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)
    ws.freeze_panes = "A2"
    wb.properties.title = sheet
    out = csv_path.with_suffix(".xlsx")
    wb.save(out)
    return out


for name in ("employee-handbook-excerpt", "benefits-summary", "onboarding-checklist", "job-description-sample"):
    print("docx:", md_to_docx(SRC / (name + ".md")).name)
print("xlsx:", csv_to_xlsx(SRC / "employee-roster-sample.csv", "Roster").name)
print("xlsx:", csv_to_xlsx(SRC / "hr-tickets-sample.csv", "Tickets").name)

ZIP = SRC.parent / "zava-sample-knowledge.zip"
with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
    for f in sorted(list(SRC.glob("*.docx")) + list(SRC.glob("*.xlsx"))):
        z.write(f, f.name)
        print("zip:", f.name)
print("saved", ZIP.name)
