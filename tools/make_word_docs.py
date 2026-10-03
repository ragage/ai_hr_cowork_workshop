"""Create Word (.docx) versions of every Markdown guide in the workshop kit.

- Uses Pandoc (via pypandoc_binary) with a custom reference.docx for consistent styling.
- Rewrites links between guides from .md to .docx so the Word set links to itself.
- Post-processes each document: table borders + header shading, footer with page numbers,
  document properties. Long guides get a table of contents (populated via Word).
"""
import pathlib
import re
import subprocess

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TOOLS = pathlib.Path(__file__).resolve().parent
ROOT = TOOLS.parent  # repo root (this file is in tools/)
BUILD = ROOT / ".build"
KIT = "Getting Things Done with Copilot Cowork for HR Tasks"
PLUM, INK, BLUE, MUTED = "3B1041", "242424", "0F6CBD", "616161"
# Prompt colour coding (callouts.lua maps <span class="goal|source|expect|constraint"> to these styles).
PROMPT_STYLES = {"Prompt Goal": ("DCEBFF", "1F4E99"), "Prompt Source": ("DFF5E1", "1E6B32"),
                 "Prompt Expectations": ("FFE9CC", "8A4B00"), "Prompt Constraints": ("EDE3FA", "5B2C91")}
TOC_DOCS = {"participant-workbook.md", "facilitator-guide.md", "facilitator-answer-key.md",
            "readiness-checklist.md", "README.md"}

SOURCES = sorted(p for p in ROOT.rglob("*.md")
                 if not any(part.startswith(".") for part in p.relative_to(ROOT).parts)  # .build, .venv, .git
                 and not {"sample-knowledge", "summary", "tools"} & set(p.parts))
CONVERTED = {p.resolve() for p in SOURCES if p.name != "SKILL.md"}


def docx_path(src):
    # Keep the Word copy of a skill outside its folder so it never ships as a companion file.
    if src.name == "SKILL.md":
        return src.parent.parent / (src.parent.name + ".SKILL.docx")
    return src.with_suffix(".docx")


# ------------------------------------------------------------------ reference document
def set_font(style, name=None, size=None, color=None, bold=None, italic=None):
    f = style.font
    if name:
        f.name = name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(a), name)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(a)) is not None:
                del rfonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic


def para_border_shade(style, left_color, fill):
    ppr = style.element.get_or_add_pPr()
    for tag in ("w:pBdr", "w:shd"):
        old = ppr.find(qn(tag))
        if old is not None:
            ppr.remove(old)
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "8"), ("w:color", left_color)):
        left.set(qn(k), v)
    bdr.append(left)
    ppr.append(bdr)
    shd = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        shd.set(qn(k), v)
    ppr.append(shd)


def make_reference():
    ref = BUILD / "reference.docx"
    data = subprocess.run([pypandoc.get_pandoc_path(), "--print-default-data-file", "reference.docx"],
                          capture_output=True, check=True).stdout
    ref.write_bytes(data)
    d = Document(ref)
    st = d.styles
    names = {s.name for s in st}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in names:
            set_font(st[name], "Segoe UI", 10.5, INK)
    st["Body Text"].paragraph_format.space_after = Pt(6)
    for lvl, size, sb in ((1, 20, 12), (2, 15, 14), (3, 12.5, 10), (4, 11, 8), (5, 10.5, 6)):
        h = st["Heading %d" % lvl]
        set_font(h, "Segoe UI Semibold", size, PLUM if lvl <= 3 else INK, bold=False, italic=False)
        h.paragraph_format.space_before = Pt(sb)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
    set_font(st["Title"], "Segoe UI Semibold", 24, PLUM)
    bt = st["Block Text"]
    set_font(bt, "Segoe UI", 10.5, INK, italic=False)
    bt.paragraph_format.left_indent = Inches(0.2)
    bt.paragraph_format.right_indent = Inches(0.1)
    bt.paragraph_format.space_before = Pt(2)
    bt.paragraph_format.space_after = Pt(2)
    para_border_shade(bt, PLUM, "FBF6FD")
    from docx.enum.style import WD_STYLE_TYPE
    for sname in ("Callout Marker",):
        if sname not in names:
            st.add_style(sname, WD_STYLE_TYPE.PARAGRAPH)
    cm = st["Callout Marker"]
    cm.base_style = st["Normal"]
    cm.font.size = Pt(2)
    cm.paragraph_format.space_before = Pt(0)
    cm.paragraph_format.space_after = Pt(0)
    for sname, (fill, color) in PROMPT_STYLES.items():
        cs = st[sname] if sname in names else st.add_style(sname, WD_STYLE_TYPE.CHARACTER)
        cs.font.color.rgb = RGBColor.from_string(color)
        cs.element.get_or_add_rPr().append(mk("w:shd", val="clear", color="auto", fill=fill))
    set_font(st["Verbatim Char"], "Consolas", 9.5, "5B2A63")
    set_font(st["Hyperlink"], color=BLUE)
    if "Source Code" in names:
        set_font(st["Source Code"], "Consolas", 9.5, INK)
    if "TOC Heading" in names:
        set_font(st["TOC Heading"], "Segoe UI Semibold", 15, PLUM, bold=False)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    d.save(ref)
    return ref


# ------------------------------------------------------------------ markdown preprocessing
LINK = re.compile(r"\]\(([^)\s]+)\)")


def preprocess(src: pathlib.Path) -> str:
    text = src.read_text(encoding="utf-8")

    # Collapsible answers: keep the content, turn the summary into a heading.
    text = re.sub(r"<details>\s*<summary>\s*(?:<strong>)?(.*?)(?:</strong>)?\s*</summary>",
                  r"#### \1", text, flags=re.S)
    text = text.replace("</details>", "")

    # SKILL.md front matter -> a readable properties table.
    if src.name == "SKILL.md" and text.startswith("---"):
        _, fm, body = text.split("---", 2)
        props = [line.split(":", 1) for line in fm.strip().splitlines() if ":" in line]
        table = ("> **Word copy for reading and printing.** Cowork uses the Markdown file "
                 "(`hr-policy-answer/SKILL.md`), not this document.\n\n"
                 "| Property | Value |\n| --- | --- |\n" +
                 "".join("| %s | %s |\n" % (k.strip().title(), v.strip()) for k, v in props) + "\n")
        body = body.lstrip()
        first_nl = body.index("\n")
        text = body[:first_nl + 1] + "\n" + table + body[first_nl + 1:]

    # Links between guides point at the Word versions (target and visible text).
    # Explicit ".md" format links (e.g. the README file map) are left pointing at the Markdown.
    def fix_link(m):
        label, link = m.group(1), m.group(2)
        if link.startswith(("http", "mailto", "#")) or label.strip() == ".md":
            return m.group(0)
        path, _, _anchor = link.partition("#")
        if not path.endswith(".md") or (src.parent / path).resolve() not in CONVERTED:
            return m.group(0)
        new_link = path[:-3] + ".docx"  # cross-document anchors aren't portable in Word
        new_label = re.sub(r"(?<![\w-])([\w./-]+)\.md\b", r"\1.docx", label)
        return "[%s](%s)" % (new_label, new_link)
    text = re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)", fix_link, text)

    # Bare mentions of kit guides in prose (e.g. "see facilitator-guide.md") -> .docx.
    # Skips link targets "(...)"; resolves the mention against this file's folder, the kit root, and the
    # audience folders the guides live in.
    def fix_bare(m):
        token = m.group(1)
        for base in (src.parent, ROOT, ROOT / "instructor", ROOT / "participant"):
            if (base / token).resolve() in CONVERTED:
                return token[:-3] + ".docx"
        return token
    text = re.sub(r"(?<![\w/.(-])((?:[\w-]+/)*[\w.-]+\.md)\b(?!\))", fix_bare, text)

    # Consecutive quotes and blank lines inside quotes are handled by the callouts.lua filter.
    return text


# ------------------------------------------------------------------ post-processing
def border_el(tag, color="D1D1D1", sz="4"):
    e = OxmlElement(tag)
    for k, v in (("w:val", "single"), ("w:sz", sz), ("w:space", "0"), ("w:color", color)):
        e.set(qn(k), v)
    return e


def style_tables(doc):
    for t in doc.tables:
        tblPr = t._tbl.tblPr
        for tag in ("w:tblBorders", "w:tblCellMar"):
            old = tblPr.find(qn(tag))
            if old is not None:
                tblPr.remove(old)
        b = OxmlElement("w:tblBorders")
        for side in ("w:top", "w:left", "w:bottom", "w:right", "w:insideH", "w:insideV"):
            b.append(border_el(side))
        tblPr.append(b)
        mar = OxmlElement("w:tblCellMar")
        for side, w in (("w:top", "40"), ("w:left", "90"), ("w:bottom", "40"), ("w:right", "90")):
            e = OxmlElement(side); e.set(qn("w:w"), w); e.set(qn("w:type"), "dxa"); mar.append(e)
        tblPr.append(mar)
        for cell in t.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "F3E5FD")):
                shd.set(qn(k), v)
            tcPr.append(shd)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(9.5)


PPR_ORDER = ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl",
             "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens",
             "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
             "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
             "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
             "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"]
RANK = {qn(t): i for i, t in enumerate(PPR_ORDER)}


def put(pPr, el):
    """Insert or replace a pPr child, keeping schema order."""
    old = pPr.find(el.tag)
    if old is not None:
        pPr.replace(old, el)
        return el
    for child in pPr:
        if RANK.get(child.tag, 99) > RANK[el.tag]:
            child.addprevious(el)
            return el
    pPr.append(el)
    return el


def mk(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), v)
    return e


def apply_callouts(doc):
    """Box every paragraph between CALLOUT markers (including list items); drop the markers."""
    body = doc.element.body
    inside = False
    for p in list(body.iterchildren(qn("w:p"))):
        txt = "".join(t.text or "" for t in p.iter(qn("w:t")))
        if txt in ("@@CALLOUT_START@@", "@@CALLOUT_END@@", "@@GAP@@"):
            if txt == "@@GAP@@":
                for t in p.iter(qn("w:t")):
                    t.text = ""
                put(p.find(qn("w:pPr")), mk("w:spacing", before="0", after="0", line="80", lineRule="exact"))
            else:
                inside = txt == "@@CALLOUT_START@@"
                body.remove(p)
            continue
        if not inside:
            continue
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr")
            p.insert(0, pPr)
        bdr = OxmlElement("w:pBdr")
        bdr.append(mk("w:left", val="single", sz="18", space="8", color=PLUM))
        put(pPr, bdr)
        put(pPr, mk("w:shd", val="clear", color="auto", fill="FBF6FD"))
        put(pPr, mk("w:spacing", before="30", after="30"))
        if pPr.find(qn("w:numPr")) is None:  # list items keep their numbering indent
            put(pPr, mk("w:ind", left="288", right="144"))


def add_field(run, instr):
    for tag, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                             ("w:instrText", {"xml:space": "preserve"}, instr),
                             ("w:fldChar", {"w:fldCharType": "end"}, None)):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        if text:
            e.text = text
        run._r.append(e)


def add_footer(doc, title):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("%s  \u00b7  %s  \u00b7  page " % (KIT, title))
        r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(MUTED); r.font.name = "Segoe UI"
        r2 = p.add_run(); r2.font.size = Pt(8); r2.font.color.rgb = RGBColor.from_string(MUTED)
        add_field(r2, "PAGE")


def first_heading(md_text):
    m = re.search(r"^# (.+)$", md_text, re.M)
    return re.sub(r"[*`]", "", m.group(1)).strip() if m else KIT


# ------------------------------------------------------------------ main
def main():
    BUILD.mkdir(exist_ok=True)
    ref = make_reference()
    made = []
    for src in SOURCES:
        md = preprocess(src)
        out = docx_path(src)
        args = ["--reference-doc", str(ref), "--lua-filter", str(TOOLS / "callouts.lua"),
                f"--resource-path={src.parent}"]
        if src.name in TOC_DOCS:
            args += ["--toc", "--toc-depth=2", "-M", "toc-title=Contents"]
        pypandoc.convert_text(md, "docx", format="gfm", outputfile=str(out), extra_args=args)
        doc = Document(out)
        title = first_heading(src.read_text(encoding="utf-8"))
        style_tables(doc)
        apply_callouts(doc)
        add_footer(doc, title)
        doc.core_properties.title = title
        doc.core_properties.subject = KIT
        doc.core_properties.author = "Copilot Cowork HR workshop kit"
        doc.save(out)
        made.append(out)
        print("docx:", out.relative_to(ROOT))
    (BUILD / "made.txt").write_text("\n".join(str(m) for m in made), encoding="utf-8")


if __name__ == "__main__":
    main()
