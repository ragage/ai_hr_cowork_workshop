"""Build a 2-page printable quick-reference card (.docx) for attendees."""
import pathlib

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = str(pathlib.Path(__file__).resolve().parents[1] / "participant" / "quick-reference-card.docx")
PLUM = RGBColor(0x3B, 0x10, 0x41)
INK = RGBColor(0x24, 0x24, 0x24)
MUTED = RGBColor(0x61, 0x61, 0x61)

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(0.5))
sec.top_margin = sec.bottom_margin = Inches(0.45)
st = doc.styles["Normal"]
st.font.name = "Segoe UI"; st.font.size = Pt(9); st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(2)

# Two-column layout
cols = sec._sectPr.xpath("./w:cols")
c = cols[0] if cols else OxmlElement("w:cols")
c.set(qn("w:num"), "2"); c.set(qn("w:space"), "360")
if not cols:
    sec._sectPr.append(c)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def h(text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = PLUM
    return p


def para(runs, style=None, size=9):
    p = doc.add_paragraph(style=style)
    for t, b in runs:
        r = p.add_run(t); r.bold = b; r.font.size = Pt(size)
    return p


# Prompt elements, same colors as the Word workbook: (shading, text color)
PROMPT_COLORS = {"G": ("DCEBFF", "1F4E99"), "S": ("DFF5E1", "1E6B32"), "E": ("FFE9CC", "8A4B00"),
                 "C": ("EDE3FA", "5B2C91")}


def cpara(runs, size=9):
    """runs: (text, element key or None, bold); element runs get their color and shading."""
    p = doc.add_paragraph()
    for t, k, b in runs:
        r = p.add_run(t); r.bold = b; r.font.size = Pt(size)
        if k:
            fill, color = PROMPT_COLORS[k]
            r.font.color.rgb = RGBColor.from_string(color)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
            r._r.get_or_add_rPr().append(shd)
    return p


def bullet(text_bold, text=""):
    return para([(text_bold, True), (text, False)], style="List Bullet")


def table(rows, widths, header_fill="F3E5FD", size=8.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j); cell.width = widths[j]
            cell.paragraphs[0].text = ""
            r = cell.paragraphs[0].add_run(val); r.font.size = Pt(size); r.bold = (i == 0)
            if i == 0:
                shade(cell, header_fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


# Title
tp = doc.add_paragraph()
r = tp.add_run("Copilot Cowork — Quick Reference Card"); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = PLUM
sp = doc.add_paragraph()
r = sp.add_run("Getting Things Done with Copilot Cowork for HR Tasks · keep this beside your laptop"); r.font.size = Pt(9); r.italic = True; r.font.color.rgb = MUTED

h("1 · Pick the right tool")
table([["Tool", "Use it when…", "HR example"],
       ["Copilot Chat", "You need a quick answer or a rewrite", "“Summarize this policy in 3 bullets.”"],
       ["Copilot Cowork", "Several steps that end in a doc, email, deck, report, or schedule", "“Draft welcome emails for 4 new hires; save as drafts.”"],
       ["Agent Builder", "Others need a reusable Q&A helper", "An HR Policy Agent employees can chat with"]],
      [Inches(1.0), Inches(2.1), Inches(1.9)])

h("2 · Write a good prompt")
cpara([("Goal", "G", True), (" + ", None, False), ("Source", "S", True), (" + ", None, False),
       ("Expectations", "E", True), (" + ", None, False), ("Constraints", "C", True)])
cpara([("e.g. “", None, False), ("Draft a warm welcome email for our new HR Coordinator starting Monday,", "G", False),
       (" ", None, False), ("using onboarding-checklist.docx.", "S", False), (" ", None, False),
       ("Under 200 words. Save as a draft", "E", False), (" — ", None, False), ("don’t send.", "C", False),
       ("”", None, False)], size=8.5)
bullet("Attach files: ", "+ → Attach cloud files, or type / and pick the file.")
bullet("Web tasks: ", "“Use Deep Research…” for cited research; “Open {site} in my browser…” runs in your Edge.")
bullet("Automations: ", "name the exact file and folder; choose Activate and run now to test.")
bullet("Always add: ", "“Save as a draft for me to review — don’t send.”")

h("3 · Approvals — one at a time")
table([["Option", "What it does", "Today?"],
       ["Send / Post / Create", "Approves this one action", "✔ Yes"],
       ["Cancel", "Skips it; Cowork continues", "✔ Yes"],
       ["Show parameters", "Shows recipients and details", "✔ Yes"],
       ["Approve All (n)", "Approves every pending action", "✘ No"],
       ["More options → Always allow", "Stops asking for this session", "✘ No"]],
      [Inches(1.7), Inches(2.4), Inches(0.9)])
para([("Clicked one by mistake? ", True), ("Side panel → Permissions → revoke.", False)])

h("4 · Where things are in Cowork")
table([["Go to…", "To…"],
       ["New task", "Start a task · model picker · reasoning effort · attach · starters"],
       ["My tasks", "Find and resume previous work"],
       ["Automations", "Schedules: Runs tab · Manage schedules (edit, pause, delete)"],
       ["Customize", "Custom instructions · your skills (Share / Re-share) · plugins"],
       ["Side panel", "Progress · Input/Output folder · skill chips · schedule · Permissions"],
       ["Output folder", "Preview · Download · Download All (zip); also in OneDrive → Cowork"]],
      [Inches(1.1), Inches(3.9)])

h("5 · Golden rules for HR")
for b, t in (("Every output is a draft. ", "Review before you send, share, or file."),
             ("Approve one at a time. ", "Read the preview before you confirm."),
             ("Fictional data in class. ", "No real employee PII."),
             ("Cite and confirm. ", "Name the source; note “please confirm with HR.”"),
             ("People decisions stay with people. ", "Cowork only sees what you’re allowed to see.")):
    bullet(b, t)

h("6 · Files: which formats work where")
table([["Where", "Use"],
       ["Today’s sample files", "Word (.docx) and Excel (.xlsx) — work in both tools"],
       ["Cowork", "Word, Excel, PowerPoint, PDF, plus .md, .csv, images, and more"],
       ["Agent Builder", ".docx, .pdf, .pptx, .txt, .xlsx — not .md or .csv"]],
      [Inches(1.9), Inches(3.1)])

h("7 · Before you leave today")
for t in ("Pause or delete the schedules from Exercises 1 and 7 (Automations → Manage schedules).",
          "Keep your custom skills set to “Only you” unless you meant to share them.",
          "Pick one task to hand to Cowork next week."):
    para([("☐ ", True), (t, False)])

doc.core_properties.title = "Copilot Cowork — Quick Reference Card"
doc.core_properties.author = "Workshop kit"
doc.save(OUT)
print("saved", OUT)
