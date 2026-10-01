#!/usr/bin/env python
"""Build the customer preview deck: walks a prospective customer through the workshop and collects
feedback before delivery. Standalone (no template needed); exercise text comes from content.py.

Environment variables (all optional):
  PREVIEW_CUSTOMER  customer name on the title slide (default "[Customer name]")
  PREVIEW_DATE      meeting date on the title slide (default "[Date]")
  PREVIEW_OUT       output path (default customer-preview-deck.pptx in the repo root)
  KIT_BUILD         scratch folder (default .build in the repo root)
"""
import os
import re

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from content import EXERCISES

TOOLS = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(TOOLS)
BUILD = os.environ.get("KIT_BUILD", os.path.join(ROOT, ".build"))
OUT = os.environ.get("PREVIEW_OUT", os.path.join(ROOT, "customer-preview-deck.pptx"))
CUSTOMER = os.environ.get("PREVIEW_CUSTOMER", "[Customer name]")
DATE = os.environ.get("PREVIEW_DATE", "[Date]")
KIT = "Getting Things Done with Copilot Cowork for HR Tasks"
os.makedirs(BUILD, exist_ok=True)

INK, SUB, LINE, WHITE = "242424", "616161", "D1D1D1", "FFFFFF"
PLUM, PURPLE, PILL, PANEL, CARD = "341434", "8064A2", "F3E5FD", "FFF9FD", "FFFCFD"
BLUE, GREEN, RED, ORANGE, TEAL = "0F6CBD", "107C10", "C0504D", "C55A11", "0E7C86"
UI, DISP, SEMI = "Segoe UI", "Segoe Sans Display", "Segoe Sans Display Semibold"
SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
TITLE_ONLY = prs.slide_layouts[5]
A16 = "http://schemas.microsoft.com/office/drawing/2014/main"
ADEC = "http://schemas.microsoft.com/office/drawing/2017/decorative"


# ---------------------------------------------------------------- background
def make_background():
    """Soft pastel wash (lavender top-left, blush top-right) on white, matching the instructor deck."""
    path = os.path.join(BUILD, "preview-bg.png")
    w, h = 160, 90
    img = Image.new("RGB", (w, h))
    px = img.load()
    blobs = [((0.05, 0.0), (236, 226, 250), 0.55), ((0.85, 0.0), (250, 226, 232), 0.6),
             ((1.0, 1.0), (240, 236, 252), 0.45)]
    for y in range(h):
        for x in range(w):
            c = [255.0, 255.0, 255.0]
            for (bx, by), col, r in blobs:
                d = ((x / w - bx) ** 2 + ((y / h - by) * 0.56) ** 2) ** 0.5
                f = max(0.0, 1 - d / r) ** 1.6
                c = [c[k] * (1 - f) + col[k] * f for k in range(3)]
            px[x, y] = tuple(int(v) for v in c)
    img.resize((1600, 900), Image.BICUBIC).save(path)
    return path


BG = make_background()


def mark_decorative(cNvPr):
    ext_lst = cNvPr.find(qn("a:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(cNvPr, qn("a:extLst"))
    ext = etree.SubElement(ext_lst, qn("a:ext"), uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}")
    etree.SubElement(ext, "{%s}decorative" % ADEC, val="1")


# ---------------------------------------------------------------- drawing helpers
def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(0.75)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if radius is not None:
        sp.adjustments[0] = radius
    mark_decorative(sp._element.nvSpPr.cNvPr)
    return sp


def rich(md, size, color, font=UI, bold_font=None):
    """'**bold** normal' -> run list."""
    runs = []
    for i, part in enumerate(md.split("**")):
        if part:
            o = {"size": size, "color": color, "font": bold_font if (i % 2 and bold_font) else font}
            if i % 2:
                o["bold"] = True
            runs.append((part, o))
    return runs


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, font=UI, color=INK,
         size=12, space_after=4, line_spacing=None):
    """paras: list of str, or dict(runs=[(text, {opts})], bullet='num'|'dot', align=...)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        if isinstance(para, str):
            para = {"runs": rich(para, size, color, font)}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        p.space_after = Pt(para.get("space_after", space_after))
        if line_spacing:
            p.line_spacing = line_spacing
        for t, o in para["runs"]:
            r = p.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = RGBColor.from_string(o.get("color", color))
            if o.get("link"):
                r.hyperlink.address = o["link"]
        bullet = para.get("bullet")
        if bullet:
            pPr = p._p.get_or_add_pPr()
            ind = Emu(Inches(0.26 if bullet == "num" else 0.2))
            pPr.set("marL", str(ind))
            pPr.set("indent", str(-ind))
            if bullet == "num":
                etree.SubElement(pPr, qn("a:buFont"), typeface="+mj-lt")
                etree.SubElement(pPr, qn("a:buAutoNum"), type="arabicPeriod")
            else:
                etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
                etree.SubElement(pPr, qn("a:buChar"), char="\u2022")
    return tb


def bullets(items, size=13, color=INK, kind="dot", font=UI):
    return [{"runs": rich(i, size, color, font), "bullet": kind} for i in items]


def chip(slide, x, y, label, fill=PILL, color=INK, size=11, h=0.32, w=None):
    w = w or 0.3 + 0.075 * len(label) * size / 11
    box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(slide, x, y, w, h, [{"runs": [(label, {"font": SEMI, "size": size, "color": color})],
                              "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    return w


def picture(slide, path, x, y, w, alt):
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


SLIDE_NO = [0]


def new_slide(title, subtitle=None, kicker="Training preview"):
    """Title-only slide with background, kicker pill, title, optional subtitle, and footer."""
    s = prs.slides.add_slide(TITLE_ONLY)
    SLIDE_NO[0] += 1
    bg = s.shapes.add_picture(BG, 0, 0, width=prs.slide_width, height=prs.slide_height)
    mark_decorative(bg._element.nvPicPr.cNvPr)
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.72), Inches(12.1), Inches(0.75)
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.text = title
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.name, r.font.size, r.font.bold = DISP, Pt(30), False
        r.font.color.rgb = RGBColor.from_string(INK)
    if kicker:
        chip(s, 0.6, 0.3, kicker, size=10, h=0.28)
    if subtitle:
        text(s, 0.6, 1.42, 12.1, 0.4, [{"runs": [(subtitle, {"size": 15, "color": SUB, "font": DISP})]}])
    text(s, 0.6, 7.05, 10.0, 0.25, [{"runs": [(KIT + "  \u00b7  Customer preview", {"size": 9, "color": SUB})]}])
    text(s, 11.9, 7.05, 0.83, 0.25, [{"runs": [(str(SLIDE_NO[0]), {"size": 9, "color": SUB})],
                                      "align": PP_ALIGN.RIGHT}])
    return s


def card(s, x, y, w, h, accent, heading, body, hsize=15, bsize=12, sub=None):
    box(s, x, y, w, h, WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    box(s, x + 0.2, y + 0.22, 0.06, 0.5 if sub else 0.38, accent)
    text(s, x + 0.38, y + 0.18, w - 0.55, 0.4, [{"runs": [(heading, {"size": hsize, "bold": True, "color": accent,
                                                                    "font": SEMI})]}])
    top = y + 0.62
    if sub:
        text(s, x + 0.38, y + 0.52, w - 0.55, 0.3, [{"runs": [(sub, {"size": 10.5, "color": SUB})]}])
        top = y + 0.9
    text(s, x + 0.25, top, w - 0.45, h - (top - y) - 0.15,
         body if isinstance(body, list) else [{"runs": rich(body, bsize, INK)}], space_after=5, line_spacing=1.05)


def band(s, y, md, fill=PLUM, color=WHITE, h=0.75, size=14):
    box(s, 0.6, y, 12.13, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    text(s, 0.9, y, 11.53, h, [{"runs": rich(md, size, color, DISP, SEMI)}], anchor=MSO_ANCHOR.MIDDLE)


def art(n):
    return os.path.join(TOOLS, "art", f"ex{n}.png")


def ex_clock(ex):
    m = re.search(r"\((\d:\d\d)-(\d:\d\d), (\d+) min\)", ex["notes_card"])
    return f"{m.group(1)}\u2013{m.group(2)}" if m else ""


# ================================================================ 1 — Title
s = prs.slides.add_slide(TITLE_ONLY)
SLIDE_NO[0] += 1
bg = s.shapes.add_picture(BG, 0, 0, width=prs.slide_width, height=prs.slide_height)
mark_decorative(bg._element.nvPicPr.cNvPr)
s.shapes._spTree.remove(bg._element)
s.shapes._spTree.insert(2, bg._element)
chip(s, 0.8, 1.35, "Training preview \u00b7 for your feedback", size=12, h=0.36)
t = s.shapes.title
t.left, t.top, t.width, t.height = Inches(0.8), Inches(1.9), Inches(8.2), Inches(2.0)
tf = t.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.TOP
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
tf.text = KIT
for r in tf.paragraphs[0].runs:
    r.font.name, r.font.size = DISP, Pt(40)
    r.font.color.rgb = RGBColor.from_string(INK)
tf.paragraphs[0].alignment = PP_ALIGN.LEFT
text(s, 0.8, 4.05, 8.0, 0.8, [{"runs": [("A 4-hour, hands-on, no-code workshop that helps HR teams hand everyday "
                                         "work to Microsoft Copilot Cowork", {"size": 17, "color": SUB, "font": DISP})]}],
     line_spacing=1.05)
box(s, 0.8, 5.15, 0.06, 0.75, PURPLE)
text(s, 1.05, 5.12, 7.5, 0.85, [{"runs": [("Prepared for ", {"size": 14, "color": SUB}),
                                          (CUSTOMER, {"size": 14, "bold": True, "color": INK})]},
                                {"runs": [(DATE, {"size": 14, "color": SUB})]}])
for i, (n, cx, cy, d) in enumerate(((1, 10.55, 1.55, 1.9), (4, 9.35, 3.55, 1.55), (8, 11.25, 4.15, 1.6))):
    box(s, cx, cy, d, d, WHITE, line=LINE, shape=MSO_SHAPE.OVAL)
    pic = picture(s, art(n), cx + d * 0.18, cy + d * 0.18, d * 0.64, "")
    mark_decorative(pic._element.nvPicPr.cNvPr)
notes(s, "TITLE. SAY: 'Thanks for your time. Today we'll walk you through the workshop we're proposing for your HR "
         "team and, more importantly, hear what you'd change. Nothing here is final.' Before the meeting, set the "
         "customer name and date on this slide (or rebuild with PREVIEW_CUSTOMER and PREVIEW_DATE).")

# ================================================================ 2 — Why we're meeting
s = new_slide("Why we\u2019re meeting today", "Preview the training, hear what fits your HR team, and shape it together")
for i, (h, body, acc) in enumerate((
        ("1 \u00b7 Preview", "Walk through the goals, agenda, and the eight hands-on exercises your team would do.", BLUE),
        ("2 \u00b7 Feedback", "Tell us which scenarios matter most, what to drop or deepen, and any data or "
                               "compliance constraints.", PURPLE),
        ("3 \u00b7 Tailor", "We adjust scenarios, sample content, timing, and format, then confirm a final plan "
                             "with you.", GREEN))):
    card(s, 0.6 + i * 4.13, 2.1, 3.87, 2.4, acc, h, body, bsize=16)
band(s, 4.95, "**Nothing is final yet.** Every exercise, timing, and example in this deck can change based on your input.",
     h=0.9, size=16)
notes(s, "WHY WE'RE MEETING. Set expectations: about 45 minutes, half walkthrough, half your feedback. The last slides "
         "are questions and a capture table; we'll fill it in together.")

# ================================================================ 3 — At a glance
s = new_slide("The workshop at a glance", "Instructor-led, hands-on, and built for HR practitioners")
for i, (big, small, acc) in enumerate((("4 hours", "one session, two breaks", BLUE), ("~25", "attendees per session", PURPLE),
                                       ("8", "hands-on exercises", GREEN), ("0", "lines of code", ORANGE))):
    x = 0.6 + i * 3.08
    box(s, x, 2.0, 2.85, 1.55, WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(s, x, 2.12, 2.85, 0.85, [{"runs": [(big, {"size": 40, "color": acc, "font": DISP})], "align": PP_ALIGN.CENTER}],
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, 2.98, 2.85, 0.4, [{"runs": [(small, {"size": 13, "color": SUB})], "align": PP_ALIGN.CENTER}])
card(s, 0.6, 3.85, 6.0, 2.95, PURPLE, "Who it\u2019s for", bullets([
    "**HR practitioners** doing day-to-day operations",
    "**Beginner-friendly:** no prior Cowork experience needed",
    "Each person works in **their own account** on a laptop with Microsoft Edge",
    "One facilitator plus **1\u20132 proctors** for a group of ~25"], size=13))
card(s, 6.73, 3.85, 6.0, 2.95, BLUE, "How it works", bullets([
    "Every exercise: **scenario card \u2192 recorded demo \u2192 hands-on**",
    "Uses **fictional sample data** (the company \u201cZava\u201d), never real employee data",
    "Attendees keep a **workbook**, a **quick-reference card**, and a **prompt library**",
    "Ends with a **30-day adoption plan**"], size=13))
notes(s, "AT A GLANCE. Ask: is ~25 the right group size? Do attendees have laptops with Edge? Mention we can run "
         "two smaller groups if needed.")

# ================================================================ 4 — Objectives
s = new_slide("What attendees will be able to do", "Six outcomes, all practiced hands-on")
OBJ = [("Choose the right tool", "Copilot Chat for quick answers, Cowork for multi-step work, Agent Builder for a "
                                 "helper others use.", BLUE),
       ("Delegate safely", "Describe an outcome, follow Cowork\u2019s plan, and approve actions one at a time.", PURPLE),
       ("Ground work in content", "Use files, cited web research, and browsing, and check outputs against sources.", GREEN),
       ("Package repeatable work", "Build and evaluate a custom skill, then schedule recurring work.", ORANGE),
       ("Build a no-code agent", "An HR policy agent grounded in documents that declines out-of-scope questions.", TEAL),
       ("Apply HR guardrails", "Permitted data only, cite and confirm policy, keep people decisions with people.", RED)]
for i, (h, b, acc) in enumerate(OBJ):
    card(s, 0.6 + (i % 3) * 4.13, 2.05 + (i // 3) * 2.45, 3.87, 2.2, acc, f"{i + 1} \u00b7 {h}", b, hsize=14, bsize=13)
notes(s, "OBJECTIVES. Ask: which of these matters most for your team? Is anything missing, e.g., a specific HR "
         "process or a manager-facing outcome?")

# ================================================================ 5 — Three tools
s = new_slide("Three Copilot tools, one workshop", "Attendees learn when to use which")
TOOLS3 = [("Copilot Chat", "Quick answers and rewrites", "\u201cSummarize this policy in 3 bullets.\u201d", BLUE,
           "Opening segment"),
          ("Copilot Cowork", "Multi-step work that ends in a document, email, deck, report, or schedule",
           "\u201cBuild a new-hire orientation deck and schedule the welcome meeting.\u201d", PURPLE, "Exercises 1\u20137"),
          ("Copilot Agent Builder", "A reusable helper other people can chat with",
           "An HR Policy Agent that answers benefits and PTO questions with sources", GREEN, "Exercise 8 (not Cowork)")]
for i, (name, use, eg, acc, where) in enumerate(TOOLS3):
    x = 0.6 + i * 4.13
    card(s, x, 2.05, 3.87, 3.3, acc, name, [{"runs": rich("**Use it when:** " + use, 15, INK)},
                                            {"runs": rich("**HR example:** " + eg, 15, INK), "space_after": 10}],
         hsize=17)
    chip(s, x + 0.25, 4.8, where, size=11.5)
band(s, 5.65, "**Cowork is the focus** (Exercises 1\u20137). The finale switches to **Agent Builder**, so attendees "
              "leave knowing when each tool fits.", h=0.85, size=15)
notes(s, "THREE TOOLS. Cowork is the focus (7 of 8 exercises). The finale is deliberately NOT Cowork: Agent Builder, "
         "so attendees leave knowing the difference between a skill that helps them and an agent that helps others.")

# ================================================================ 6 — Agenda
s = new_slide("Agenda \u00b7 240 minutes", "Short framing, then mostly hands-on")
AGENDA = [("0:00\u20130:40", "Welcome, Copilot vs. Cowork, UI tour, and setup")]
for ex in EXERCISES:
    AGENDA.append((ex_clock(ex), f"Ex {ex['num']} \u00b7 {ex['title']}" + (" (not Cowork)" if ex.get("not_cowork") else "")))
    if ex["num"] == 2:
        AGENDA.append(("1:25\u20131:35", "Break"))
    if ex["num"] == 5:
        AGENDA.append(("2:45\u20132:55", "Break"))
    if ex["num"] == 7:
        AGENDA.append(("3:30\u20133:35", "Spotlight \u00b7 Extend Cowork with HR plugins (talk only)"))
AGENDA.append(("3:55\u20134:00", "Wrap-up and next steps"))
half = (len(AGENDA) + 1) // 2
for col, rows in enumerate((AGENDA[:half], AGENDA[half:])):
    for j, (tm, seg) in enumerate(rows):
        x, y = 0.6 + col * 6.2, 2.0 + j * 0.66
        is_break = seg.startswith("Break")
        box(s, x, y, 5.9, 0.56, PANEL if not is_break else "F4F4F4", line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            radius=0.2)
        text(s, x + 0.2, y, 1.3, 0.56, [{"runs": [(tm, {"size": 12, "bold": True, "color": PURPLE, "font": SEMI})]}],
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 1.5, y, 4.3, 0.56, [{"runs": [(seg, {"size": 12.5, "color": SUB if is_break else INK})]}],
             anchor=MSO_ANCHOR.MIDDLE)
notes(s, "AGENDA. Ask: does one 4-hour block work, or would two 2-hour sessions fit your calendar better? Any hard "
         "stop times?")


# ================================================================ 7-8 — Exercises
def exercise_slide(exs, title):
    s = new_slide(title, "What attendees build, and why it matters for HR")
    w, gap = 2.9, 0.177
    for i, ex in enumerate(exs):
        x, y, h = 0.6 + i * (w + gap), 1.95, 4.9
        nc = ex.get("not_cowork")
        box(s, x, y, w, h, WHITE, line=RED if nc else LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
        box(s, x + w / 2 - 0.5, y + 0.2, 1.0, 1.0, PANEL, line=LINE, shape=MSO_SHAPE.OVAL)
        picture(s, art(ex["num"]), x + w / 2 - 0.34, y + 0.36, 0.68, f"Illustration for Exercise {ex['num']}")
        if nc:
            chip(s, x + 0.15, y + 0.15, "Not Cowork", fill=RED, color=WHITE, size=9.5, h=0.27)
        text(s, x + 0.2, y + 1.32, w - 0.4, 0.28, [{"runs": [(f"EXERCISE {ex['num']}", {"size": 10, "bold": True,
                                                                                         "color": PURPLE})],
                                                   "align": PP_ALIGN.CENTER}])
        text(s, x + 0.2, y + 1.58, w - 0.4, 0.72, [{"runs": [(ex["title"], {"size": 14.5, "bold": True, "color": INK,
                                                                             "font": SEMI})],
                                                   "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
        text(s, x + 0.22, y + 2.38, w - 0.44, 1.85, [{"runs": rich(ex["goal"], 11.5, INK)},
                                                      {"runs": rich("**Output:** " + ex["output"], 10, SUB)}],
             space_after=6, line_spacing=1.02)
        cx = x + 0.22
        cx += chip(s, cx, y + h - 0.5, ex["minutes"], size=10, h=0.3) + 0.1
        chip(s, cx, y + h - 0.5, ex["function"], size=10, h=0.3)
    notes(s, "EXERCISES. For each card, give one sentence on the HR value. ASK: which of these would your team use "
             "next week? Which would you drop or swap for one of your own HR scenarios? "
             + ("Exercise 8 is deliberately NOT Cowork: it's Copilot Agent Builder." if any(e.get("not_cowork") for e in exs) else ""))
    return s


exercise_slide(EXERCISES[:4], "The hands-on exercises \u00b7 1\u20134")
exercise_slide(EXERCISES[4:], "The hands-on exercises \u00b7 5\u20138")

# ================================================================ 9 — How each exercise runs
s = new_slide("How each exercise runs", "The same rhythm every time, so attendees always know what\u2019s next")
STEPS = [("Scenario card", "Goal, output, why Cowork, the prompt, workflow, and data sources", BLUE),
         ("Recorded demo", "The facilitator plays a recorded run, so no one waits for long tasks", PURPLE),
         ("Hands-on", "Each attendee runs it in their own account, with proctors on hand", GREEN),
         ("Checkpoint", "Check the result against the expected outcome, then discuss", ORANGE)]
for i, (h, b, acc) in enumerate(STEPS):
    x = 0.6 + i * 3.1
    box(s, x, 2.1, 2.8, 2.7, WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    box(s, x + 1.05, 2.3, 0.7, 0.7, acc, shape=MSO_SHAPE.OVAL)
    text(s, x + 1.05, 2.3, 0.7, 0.7, [{"runs": [(str(i + 1), {"size": 20, "bold": True, "color": WHITE})],
                                       "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, 3.15, 2.4, 0.4, [{"runs": [(h, {"size": 16, "bold": True, "color": acc, "font": SEMI})],
                                       "align": PP_ALIGN.CENTER}])
    text(s, x + 0.2, 3.6, 2.4, 1.1, [{"runs": [(b, {"size": 12.5, "color": INK})], "align": PP_ALIGN.CENTER}],
         line_spacing=1.05)
    if i < 3:
        text(s, x + 2.8, 3.1, 0.3, 0.5, [{"runs": [("\u2192", {"size": 22, "color": SUB})], "align": PP_ALIGN.CENTER}])
band(s, 5.2, "**Facilitators come prepared:** an answer key with expected results, a troubleshooting guide, and a "
             "whole-room Plan B.", h=0.85)
notes(s, "EXERCISE RHYTHM. Highlight the recorded demos: Cowork tasks can take several minutes, so the room watches a "
         "recording and then does it live. Ask: do you want attendees to receive the demo recordings afterward?")

# ================================================================ 10 — Safety
s = new_slide("Built-in safety for HR work", "Habits we practice in every exercise")
SAFE = [("Fictional data only", "Exercises use the fictional company Zava. No real employee data.", BLUE),
        ("Approve one at a time", "Cowork asks before it sends, posts, or schedules. No \u201cApprove All\u201d.", PURPLE),
        ("Every output is a draft", "Review before anything is sent, shared, or filed.", GREEN),
        ("Cite and confirm", "Policy answers name their source and say \u201cplease confirm with HR\u201d.", ORANGE),
        ("People decisions stay with people", "No performance ratings or pay decisions by AI.", TEAL),
        ("Your own space", "Each attendee works in their own OneDrive; skills stay private.", RED)]
for i, (h, b, acc) in enumerate(SAFE):
    card(s, 0.6 + (i % 3) * 4.13, 2.05 + (i // 3) * 2.3, 3.87, 2.05, acc, h, b, hsize=14, bsize=13)
notes(s, "SAFETY. Ask: any policies we must reflect, such as works council agreements, regional data rules, or "
         "restrictions on which HR topics can be used in training?")

# ================================================================ 11 — Takeaways
s = new_slide("What your team takes away", "Materials and habits that last beyond the session")
TAKE = [("Participant workbook", "Every exercise, step by step, with checkpoints"),
        ("Quick-reference card", "One printable page: tools, prompt recipe, approvals"),
        ("Prompt library", "Copy-paste HR prompts by scenario"),
        ("Their own custom skill", "An HR Policy Answer skill they built and tested"),
        ("A working HR agent", "Built in Agent Builder, grounded in HR documents"),
        ("After-the-workshop pack", "Knowledge check, feedback survey, and a 30-day adoption plan")]
for i, (h, b) in enumerate(TAKE):
    x, y = 0.6 + (i % 2) * 6.13, 2.05 + (i // 2) * 1.5
    box(s, x, y, 5.9, 1.3, WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    box(s, x + 0.25, y + 0.35, 0.6, 0.6, PILL, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.25, y + 0.35, 0.6, 0.6, [{"runs": [("\u2713", {"size": 18, "bold": True, "color": PURPLE})],
                                            "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.1, y + 0.2, 4.6, 0.9, [{"runs": [(h, {"size": 15, "bold": True, "color": INK, "font": SEMI})]},
                                         {"runs": [(b, {"size": 12.5, "color": SUB})]}], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "TAKEAWAYS. Ask: would you like any of these branded or translated? Should managers get a short version?")

# ================================================================ 12 — What we need from you
s = new_slide("What we need from you", "Tenant readiness is the biggest success factor")
card(s, 0.6, 2.0, 6.0, 4.85, PURPLE, "Prerequisites", bullets([
    "A **Microsoft 365 Copilot** license for each attendee",
    "**Usage-based billing** and a **Cowork spending policy** for the attendee group (this grants Cowork access)",
    "**Cowork Browsing** allowed for the group (Exercise 3)",
    "~25 attendee accounts **plus 2\u20133 spares** in one tenant",
    "Laptops with **Microsoft Edge 152 or later**",
    "A shared Teams or SharePoint folder for the sample files"], size=14.5))
TL = [("T \u2212 3 weeks", "Licenses, billing, spending policy, browsing"),
      ("T \u2212 2 weeks", "Accounts and spares provisioned"),
      ("T \u2212 1 week", "Full dry run; demo videos recorded"),
      ("T \u2212 1 day", "Sample emails and meetings loaded; spot checks"),
      ("Day after", "Pause schedules; survey; review usage")]
box(s, 6.73, 2.0, 6.0, 4.85, WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, 7.05, 2.18, 5.5, 0.4, [{"runs": [("Preparation timeline", {"size": 15, "bold": True, "color": BLUE, "font": SEMI})]}])
for j, (when, what) in enumerate(TL):
    y = 2.75 + j * 0.8
    box(s, 7.05, y + 0.12, 0.22, 0.22, BLUE, shape=MSO_SHAPE.OVAL)
    if j < len(TL) - 1:
        box(s, 7.15, y + 0.34, 0.02, 0.58, LINE)
    text(s, 7.45, y, 5.1, 0.75, [{"runs": [(when, {"size": 14.5, "bold": True, "color": INK, "font": SEMI})]},
                                 {"runs": [(what, {"size": 13.5, "color": SUB})]}], space_after=1)
notes(s, "WHAT WE NEED. Ask: who is the tenant admin contact? Are Copilot licenses already assigned? Is usage-based "
         "billing set up? These drive the date we can commit to.")

# ================================================================ 13 — Cost
s = new_slide("Planning for cost", "Cowork is usage-billed, so we plan and cap it together")
card(s, 0.6, 2.0, 6.0, 3.6, GREEN, "Before the session", bullets([
    "**Estimate** with the Copilot Cowork customer cost estimator",
    "**Cap it:** a per-user credit limit and alerts on the workshop spending policy",
    "**Keep defaults** (\u201cLet Cowork decide\u201d model, default reasoning effort) unless an exercise needs more"],
    size=15), hsize=17)
card(s, 6.73, 2.0, 6.0, 3.6, ORANGE, "During and after", bullets([
    "Attendees check what a task cost with **/cost** (taught in Exercise 1)",
    "**Pause** the two scheduled automations the day after",
    "**Remove** the workshop group from the spending policy if accounts shouldn\u2019t keep access",
    "Review totals in the **Microsoft 365 admin center**"], size=15), hsize=17)
text(s, 0.6, 5.85, 12.1, 0.4, [{"runs": [("Estimator: ", {"size": 13, "color": SUB}),
                                         ("aka.ms/CustomerCoworkEstimator",
                                          {"size": 13, "color": BLUE, "link": "https://aka.ms/CustomerCoworkEstimator"}),
                                         ("    \u00b7    Credit usage and /cost: ", {"size": 13, "color": SUB}),
                                         ("learn.microsoft.com",
                                          {"size": 13, "color": BLUE,
                                           "link": "https://learn.microsoft.com/microsoft-365/copilot/usage-based-billing-copilot-credits-cost"})]}])
notes(s, "COST. Don't quote a number live unless you've run the estimator for their size. Ask: who approves the "
         "budget, and what per-user limit feels comfortable?")

# ================================================================ 14 — Feedback questions
s = new_slide("Your feedback: help us tailor it", "Questions for today\u2019s discussion", kicker="Your input")
QS = [("Audience & goals", ["Who will attend, and how familiar are they with Copilot?",
                            "What should they be able to do the week after?"], BLUE),
      ("Scenarios", ["Which exercises matter most? What would you drop or deepen?",
                     "Which HR tasks take the most manual effort today?"], PURPLE),
      ("Data & systems", ["Which HR systems do you use (e.g., SAP SuccessFactors, Workday)?",
                          "Fictional data, or sanitized versions of your own policies?",
                          "Any data, regional, or works council constraints?"], GREEN),
      ("Format & logistics", ["One 4-hour block, or two 2-hour sessions? In person or remote?",
                              "Preferred dates, and who owns tenant readiness?"], ORANGE)]
for i, (h, qs, acc) in enumerate(QS):
    card(s, 0.6 + (i % 2) * 6.13, 2.0 + (i // 2) * 2.45, 5.9, 2.25, acc, h, bullets(qs, size=13))
notes(s, "FEEDBACK QUESTIONS. Spend most of the meeting here. Capture answers on the next slide, live, so the "
         "customer sees what we heard.")

# ================================================================ 15 — Tailoring options
s = new_slide("Ways we can tailor the workshop", "Pick what fits; we\u2019ll confirm the trade-offs on time", kicker="Your input")
OPTS = [("Swap in your scenarios", "Replace an exercise with one of your own HR workflows.", BLUE),
        ("Use your content", "Sanitized versions of your handbook or policies instead of the Zava samples.", PURPLE),
        ("Spotlight your HR system", "How your HR system could connect, e.g., an agent or a plugin for SAP "
                                     "SuccessFactors.", GREEN),
        ("Change the format", "Two 2-hour sessions, or a 90-minute overview for HR leaders.", ORANGE),
        ("Go deeper on agents", "Extend Exercise 8, or plan a follow-up on Copilot Studio.", TEAL),
        ("Follow-up support", "Office hours and a 30-day adoption check-in.", RED)]
for i, (h, b, acc) in enumerate(OPTS):
    card(s, 0.6 + (i % 3) * 4.13, 2.05 + (i // 3) * 2.3, 3.87, 2.05, acc, h, b, hsize=14, bsize=13)
notes(s, "TAILORING. Note any choices on the capture slide. Swapping exercises keeps the 4-hour length; adding "
         "content means dropping something else.")

# ================================================================ 16 — Feedback capture
s = new_slide("Feedback capture", "Filled in together during the meeting", kicker="Your input")
rows = [("Topic", "What we heard", "Priority"), ("Audience & goals", "", ""), ("Exercises to keep / change", "", ""),
        ("Data & HR systems", "", ""), ("Format & logistics", "", ""), ("Tenant readiness & owner", "", ""),
        ("Other", "", "")]
gf = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(2.0), Inches(12.13), Inches(4.8))
gf.name = "Feedback capture table"
tbl = gf.table
for j, wd in enumerate((3.0, 7.63, 1.5)):
    tbl.columns[j].width = Inches(wd)
for i, row in enumerate(rows):
    tbl.rows[i].height = Inches(0.6 if i == 0 else 0.7)
    for j, val in enumerate(row):
        c = tbl.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor.from_string(PILL if i == 0 else WHITE)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.text = val
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(13 if i == 0 else 12.5)
                r.font.bold = i == 0 or j == 0
                r.font.name = UI
                r.font.color.rgb = RGBColor.from_string(INK)
notes(s, "FEEDBACK CAPTURE. Type into the table live (Priority: High / Medium / Low). Afterward, save this deck and "
         "send it back to the customer as the meeting summary.")

# ================================================================ 17 — Next steps
s = new_slide("Next steps", "From today\u2019s feedback to a workshop that fits your team")
NEXT = [("Today", "Your feedback on scenarios, data, and format", BLUE),
        ("Within 1 week", "We share a tailored agenda and exercise list", PURPLE),
        ("T \u2212 3 weeks", "Tenant readiness starts with your admin", GREEN),
        ("T \u2212 1 week", "Full dry run in your tenant", ORANGE),
        ("Workshop day", "4 hours, hands-on", TEAL),
        ("+ 30 days", "Adoption check-in", RED)]
for i, (when, what, acc) in enumerate(NEXT):
    x = 0.6 + i * 2.05
    box(s, x, 2.95, 1.85, 0.12, acc)
    box(s, x + 0.66, 2.66, 0.54, 0.54, acc, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.66, 2.66, 0.54, 0.54, [{"runs": [(str(i + 1), {"size": 15, "bold": True, "color": WHITE})],
                                          "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, 3.45, 1.85, 0.4, [{"runs": [(when, {"size": 16, "bold": True, "color": acc, "font": SEMI})],
                                  "align": PP_ALIGN.CENTER}])
    text(s, x + 0.05, 3.95, 1.75, 1.3, [{"runs": [(what, {"size": 14, "color": INK})], "align": PP_ALIGN.CENTER}],
         line_spacing=1.05)
band(s, 5.65, "**Thank you.** Questions or more feedback? Contact: [Name] \u00b7 [Email]", h=0.9, size=17)
notes(s, "NEXT STEPS. Agree on the owner and date for each step before you leave. Fill in your contact details on "
         "this slide before the meeting.")

# ---------------------------------------------------------------- properties and save
prs.core_properties.title = KIT + " \u2014 Customer preview"
prs.core_properties.subject = "Training preview for customer feedback"
prs.core_properties.author = "Copilot Cowork HR workshop kit"
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
